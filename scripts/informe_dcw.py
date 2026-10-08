"""Estudio de feria de Data Centre World Madrid 2026 (Word, PDF y Excel).

Uso:
    python scripts/informe_dcw.py [--fecha AAAA-MM-DD]

Lee datos/estudios/dcw_madrid_2026.json, comprueba los términos prohibidos del
apartado 8 de CLAUDE.md y deja en salidas/:
    Estudio_Data_Centre_World_Madrid_2026_AAAA-MM-DD.docx y .pdf
    Contactos_Data_Centre_World_Madrid_2026_AAAA-MM-DD.xlsx
    Contactos_Data_Centre_World_Madrid_2026_AAAA-MM-DD_Apollo.csv
"""
import argparse
import csv
import datetime
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import bloc_pdf
from comun import RAIZ, SALIDAS
from validar import EMOJI, MARCADORES, PROHIBIDOS, EMAIL, es_generico

DATOS_ESTUDIO = RAIZ / 'datos' / 'estudios' / 'dcw_madrid_2026.json'

FUENTE = 'Calibri'
GRIS_CABECERA = 'D9D9D9'
GRIS_TEXTO = RGBColor(0x59, 0x59, 0x59)
PERFILES = {
    'A': 'Ingeniería MEP, mecánica, PCI o diseño de centros de datos',
    'B': 'Responsable de centros de datos o misión crítica',
    'C': 'Gestión de proyecto, compras o procurement',
    'D': 'Dirección general o de oficina',
}
PRIORIDADES = {
    1: 'Prioridad 1 · Ingenierías de diseño, MEP y misión crítica',
    2: 'Prioridad 2 · Contratistas MEP, EPC y PCI',
    3: 'Prioridad 3 · Gestión de proyectos, costes y procurement',
    4: 'Prioridad 4 · Constructoras con división de centros de datos',
}


# ----------------------------------------------------------------- comprobación

def cadenas(obj, ruta=''):
    if isinstance(obj, str):
        yield ruta, obj
    elif isinstance(obj, dict):
        for k, v in obj.items():
            yield from cadenas(v, f'{ruta}.{k}' if ruta else k)
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from cadenas(v, f'{ruta}[{i}]')


def comprobar(datos):
    """Errores (términos prohibidos, marcadores, emojis, exclamaciones) y avisos (correos generales)."""
    errores, avisos = [], []
    for ruta, texto in cadenas(datos):
        if ruta.startswith('fuentes'):
            continue
        for regla, nombre in ((PROHIBIDOS, 'término prohibido'), (MARCADORES, 'marcador no admitido'),
                              (EMOJI, 'emoji')):
            m = regla.search(texto)
            if m:
                errores.append(f'{ruta}: {nombre} «{m.group(0)}».')
        if '!' in texto or '¡' in texto:
            errores.append(f'{ruta}: signo de exclamación.')
        for m in EMAIL.finditer(texto):
            if es_generico(m.group(0)):
                avisos.append(f'{ruta}: correo general «{m.group(0)}» (se mantiene hasta tener uno personal).')
    return errores, avisos


def contar(datos):
    contactos = [c for e in datos['empresas'] for c in e['contactos']]
    ab = [c for e in datos['empresas'] if e['prioridad'] <= 2 for c in e['contactos'] if c['perfil'] in 'AB']
    return {'n_contactos': len(contactos), 'n_empresas': len(datos['empresas']), 'n_ab_p1': len(ab)}


# ----------------------------------------------------------------- Word: utilidades

def fuente_run(run, tam=None, negrita=None, color=None, cursiva=None):
    run.font.name = FUENTE
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FUENTE)
    if tam:
        run.font.size = Pt(tam)
    if negrita is not None:
        run.bold = negrita
    if cursiva is not None:
        run.italic = cursiva
    if color is not None:
        run.font.color.rgb = color
    return run


def parrafo(doc, texto='', tam=10, negrita=False, estilo=None, alinear=None, espacio_despues=4, color=None):
    p = doc.add_paragraph(style=estilo) if estilo else doc.add_paragraph()
    if texto:
        fuente_run(p.add_run(texto), tam, negrita, color)
    p.paragraph_format.space_after = Pt(espacio_despues)
    if alinear:
        p.alignment = alinear
    return p


def vineta(doc, texto, tam=10):
    p = doc.add_paragraph(style='List Bullet')
    fuente_run(p.add_run(texto), tam)
    p.paragraph_format.space_after = Pt(3)
    return p


def titulo(doc, texto, nivel=1):
    h = doc.add_heading(level=nivel)
    fuente_run(h.add_run(texto), {1: 15, 2: 12, 3: 10.5}[nivel], True, RGBColor(0, 0, 0))
    h.paragraph_format.space_before = Pt({1: 14, 2: 10, 3: 8}[nivel])
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    return h


def sombrear(celda, color):
    tcPr = celda._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def repetir_cabecera(fila):
    trPr = fila._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def no_partir(fila):
    trPr = fila._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def fijar_anchos(t, anchos):
    """Anchos fijos de columna (LibreOffice ignora los anchos de celda si la tabla no los declara)."""
    tblPr = t._tbl.tblPr
    for nombre in ('w:tblW', 'w:tblLayout'):
        for viejo in tblPr.findall(qn(nombre)):
            tblPr.remove(viejo)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(sum(anchos) / 2.54 * 1440)))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    grid = t._tbl.tblGrid
    for i, col in enumerate(grid.findall(qn('w:gridCol'))):
        col.set(qn('w:w'), str(int(anchos[i] / 2.54 * 1440)))
    for fila in t.rows:
        for i, ancho in enumerate(anchos):
            fila.cells[i].width = Cm(ancho)


def texto_celda(celda, texto, tam=8.5, negrita=False, color=None):
    celda.text = ''
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    fuente_run(p.add_run(texto or ''), tam, negrita, color)


def tabla(doc, cabecera, filas, anchos, tam=8.5, cabecera_gris=True):
    t = doc.add_table(rows=1, cols=len(cabecera))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, txt in enumerate(cabecera):
        c = t.rows[0].cells[i]
        texto_celda(c, txt, tam, True)
        if cabecera_gris:
            sombrear(c, GRIS_CABECERA)
    repetir_cabecera(t.rows[0])
    for fila in filas:
        celdas = t.add_row().cells
        for i, txt in enumerate(fila):
            texto_celda(celdas[i], txt, tam)
        no_partir(t.rows[-1])
    fijar_anchos(t, anchos)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def tabla_ficha(doc, pares, anchos=(3.4, 13.6), tam=8.5):
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Table Grid'
    t.autofit = False
    for k, v in pares:
        if not v:
            continue
        celdas = t.add_row().cells
        texto_celda(celdas[0], k, tam, True)
        sombrear(celdas[0], 'F2F2F2')
        texto_celda(celdas[1], v, tam)
        no_partir(t.rows[-1])
    fijar_anchos(t, anchos)
    return t


def campo_pagina(parrafo_pie):
    run = parrafo_pie.add_run()
    fuente_run(run, 8, color=GRIS_TEXTO)
    for tipo, texto in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tipo:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tipo)
        else:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = texto
        run._r.append(el)


def preparar_seccion(seccion, pie, apaisada=False):
    if apaisada:
        seccion.orientation = WD_ORIENT.LANDSCAPE
        seccion.page_width, seccion.page_height = Cm(29.7), Cm(21.0)
    else:
        seccion.orientation = WD_ORIENT.PORTRAIT
        seccion.page_width, seccion.page_height = Cm(21.0), Cm(29.7)
    seccion.left_margin = seccion.right_margin = Cm(2.0)
    seccion.top_margin, seccion.bottom_margin = Cm(1.8), Cm(1.8)
    seccion.footer.is_linked_to_previous = False
    p = seccion.footer.paragraphs[0]
    p.text = ''
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fuente_run(p.add_run(pie + ' · Página '), 8, color=GRIS_TEXTO)
    campo_pagina(p)


def nueva_seccion(doc, pie, apaisada):
    s = doc.add_section()
    preparar_seccion(s, pie, apaisada)
    return s


# ----------------------------------------------------------------- Word: contenido

def portada(doc, est, fecha_txt):
    for _ in range(6):
        parrafo(doc, '')
    parrafo(doc, 'STORA', 30, True, alinear=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=2)
    parrafo(doc, 'Estudio de feria', 14, alinear=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=24, color=GRIS_TEXTO)
    parrafo(doc, est['titulo'], 24, True, alinear=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=8)
    parrafo(doc, est['subtitulo'], 13, alinear=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=8)
    parrafo(doc, est['evento_linea'], 12, alinear=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=60)
    parrafo(doc, 'Alliance Albors y Prado S.L.', 11, True, alinear=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=2)
    parrafo(doc, f"Datos a {est['fecha_datos']} · Documento de {fecha_txt}", 10,
            alinear=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=2)
    parrafo(doc, 'Uso interno', 10, alinear=WD_ALIGN_PARAGRAPH.CENTER, color=GRIS_TEXTO)


CAPITULOS = (
    '1. Resumen ejecutivo',
    '2. La feria: datos, programa y expositores',
    '3. Mercado y proyectos en España y Portugal',
    '4. Quién decide el depósito en un centro de datos',
    '5. Empresas objetivo y contactos',
    '6. Operadores: interlocutores de la especificación del propietario',
    '7. Plan de acción',
    '8. Metodología y grado de certeza',
    '9. Fuentes consultadas',
)


def contenido(doc):
    titulo(doc, 'Contenido', 1)
    for c in CAPITULOS:
        parrafo(doc, c, 10.5, espacio_despues=3)


def cap_resumen(doc, d, n):
    titulo(doc, CAPITULOS[0], 1)
    for texto in d['resumen']:
        vineta(doc, texto.format(**n))


def cap_feria(doc, d):
    f = d['feria']
    titulo(doc, CAPITULOS[1], 1)
    titulo(doc, '2.1 Datos clave', 2)
    tabla_ficha(doc, f['datos'])
    parrafo(doc, '')
    titulo(doc, '2.2 Programa de interés para Stora (edición 2026)', 2)
    tabla(doc, ['Ponente', 'Cargo y empresa', 'Sesión', 'Día y hora', 'Interés'],
          [[s['ponente'], ', '.join(x for x in (s['cargo'], s['empresa']) if x), s['sesion'], s['cuando'], s['interes']]
           for s in f['programa']],
          [3.3, 4.6, 4.6, 2.8, 1.7])
    parrafo(doc, f['programa_nota'], 8.5, color=GRIS_TEXTO)
    titulo(doc, '2.3 Ponentes de ediciones anteriores', 2)
    tabla(doc, ['Ponente', 'Cargo', 'Empresa', 'Edición'],
          [[p['ponente'], p['cargo'], p['empresa'], p['edicion']] for p in f['ponentes_anteriores']],
          [3.6, 6.2, 4.6, 2.6])
    titulo(doc, '2.4 Expositores relevantes', 2)
    tabla(doc, ['Empresa', 'Actividad', 'Stand o edición'], f['expositores'], [5.5, 7.0, 4.5])
    titulo(doc, '2.5 Otros eventos donde coinciden los mismos interlocutores', 2)
    tabla(doc, ['Evento', 'Fecha y lugar', 'Participantes de interés'], f['eventos_relacionados'], [4.2, 4.0, 8.8])
    titulo(doc, '2.6 Recomendaciones para la visita', 2)
    for r in f['recomendaciones']:
        vineta(doc, r)


def cap_mercado(doc, d, pie):
    m = d['mercado']
    nueva_seccion(doc, pie, apaisada=True)
    titulo(doc, CAPITULOS[2], 1)
    for c in m['cifras']:
        vineta(doc, c)
    titulo(doc, '3.1 Proyectos y agentes conocidos', 2)
    tabla(doc, ['Promotor', 'Ubicación', 'Potencia', 'Inversión', 'Estado', 'Ingeniería y construcción conocidas', 'Interés para Stora'],
          [[p['promotor'], p['ubicacion'], p['potencia'], p['inversion'], p['estado'], p['agentes'], p['interes']]
           for p in m['proyectos']],
          [2.8, 3.6, 2.5, 2.4, 3.7, 6.4, 4.3], tam=8)
    parrafo(doc, m['proyectos_nota'], 8.5, color=GRIS_TEXTO)


def cap_decision(doc, d, pie):
    nueva_seccion(doc, pie, apaisada=False)
    titulo(doc, CAPITULOS[3], 1)
    for p in d['decision']['parrafos']:
        parrafo(doc, p, 10, espacio_despues=6)
    titulo(doc, '4.1 Marco normativo que fija la reserva de agua', 2)
    tabla(doc, ['Norma', 'Qué exige'], d['decision']['normativa'], [5.0, 12.0])


def mantener_junto(t, hasta_ultima=True):
    """Mantiene la tabla en la misma página que lo que la sigue (o, sin la última fila, unida en sí misma)."""
    filas = t.rows if hasta_ultima else t.rows[:-1]
    for fila in filas:
        for celda in fila.cells:
            for p in celda.paragraphs:
                p.paragraph_format.keep_with_next = True


def ficha_empresa(doc, e):
    titulo(doc, f"{e['nombre']}", 3)
    sub = parrafo(doc, f"Prioridad {e['prioridad']} · {e['categoria']}", 9, color=GRIS_TEXTO, espacio_despues=3)
    sub.paragraph_format.keep_with_next = True
    ficha = tabla_ficha(doc, [
        ('Papel', e['rol']),
        ('Centros de datos', e['dc']),
        ('Agua, PCI y refrigeración', e['agua_pci']),
        ('Clientes documentados', e['clientes']),
        ('Ferias y foros', e['feria']),
        ('Oficinas', e['oficinas']),
        ('Teléfono', e['telefono']),
        ('Correo general', e['email']),
        ('Web', e['web']),
        ('Por qué', e['motivo']),
    ])
    if e['contactos']:
        mantener_junto(ficha)
        parrafo(doc, '', espacio_despues=1).paragraph_format.keep_with_next = True
        filas = []
        for c in sorted(e['contactos'], key=lambda c: c['perfil']):
            via = c['linkedin'] or ''
            filas.append([c['nombre'], c['cargo'], c['perfil'], c['ubicacion'], c['origen'], via, c['certeza']])
        contactos = tabla(doc, ['Nombre', 'Cargo', 'Perfil', 'Ubicación', 'Origen del dato', 'LinkedIn', 'Certeza'],
                          filas, [2.7, 4.0, 1.25, 1.65, 3.0, 2.4, 2.0], tam=8)
        mantener_junto(contactos, hasta_ultima=False)
    else:
        mantener_junto(ficha, hasta_ultima=False)
        parrafo(doc, '', espacio_despues=6)


def cap_empresas(doc, d):
    titulo(doc, CAPITULOS[4], 1)
    parrafo(doc, 'Perfiles de contacto: ' + '; '.join(f'{k}, {v}' for k, v in PERFILES.items()) +
            '. Los perfiles A son los prioritarios para entrar en la especificación.', 9, espacio_despues=6)
    titulo(doc, '5.1 Resumen por prioridad', 2)
    filas = [[str(e['prioridad']), e['nombre'], e['categoria'], str(len(e['contactos'])), e['motivo']]
             for e in d['empresas']]
    tabla(doc, ['Prior.', 'Empresa', 'Categoría', 'Contactos', 'Por qué'], filas, [1.2, 4.2, 3.7, 1.9, 6.0])
    for prioridad, texto in PRIORIDADES.items():
        grupo = [e for e in d['empresas'] if e['prioridad'] == prioridad]
        if not grupo:
            continue
        titulo(doc, '5.' + str(prioridad + 1) + ' ' + texto, 2)
        for e in grupo:
            ficha_empresa(doc, e)
    titulo(doc, '5.6 Otras empresas a seguir', 2)
    tabla(doc, ['Empresa', 'Motivo'], d['otras'], [5.0, 12.0])


def cap_operadores(doc, d):
    titulo(doc, CAPITULOS[5], 1)
    parrafo(doc, 'El propietario fija el estándar técnico y aprueba las desviaciones. Estas personas no compran el depósito, '
                 'pero sus equipos de diseño y construcción validan las especificaciones de sus ingenierías.', 10)
    tabla(doc, ['Operador', 'Nombre', 'Cargo', 'Dónde consta'], d['operadores'], [3.4, 3.6, 6.0, 4.0])


def cap_plan(doc, d, n):
    titulo(doc, CAPITULOS[6], 1)
    for p in d['plan']:
        vineta(doc, p.format(**n))


def cap_metodologia(doc, d):
    titulo(doc, CAPITULOS[7], 1)
    for p in d['metodologia']:
        vineta(doc, p)


def cap_fuentes(doc, d):
    titulo(doc, CAPITULOS[8], 1)
    for i, (nombre, url) in enumerate(d['fuentes'], 1):
        p = parrafo(doc, '', espacio_despues=2)
        fuente_run(p.add_run(f'{i}. {nombre}. '), 8.5)
        fuente_run(p.add_run(url), 8, color=GRIS_TEXTO)


def generar_docx(d, ruta, fecha_txt, n):
    doc = Document()
    estilo = doc.styles['Normal']
    estilo.font.name = FUENTE
    estilo.element.rPr.rFonts.set(qn('w:eastAsia'), FUENTE)
    estilo.font.size = Pt(10)
    pie = d['estudio']['pie']
    preparar_seccion(doc.sections[0], pie)
    portada(doc, d['estudio'], fecha_txt)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    contenido(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    cap_resumen(doc, d, n)
    cap_feria(doc, d)
    cap_mercado(doc, d, pie)
    cap_decision(doc, d, pie)
    cap_empresas(doc, d)
    cap_operadores(doc, d)
    cap_plan(doc, d, n)
    cap_metodologia(doc, d)
    cap_fuentes(doc, d)
    doc.core_properties.title = f"{d['estudio']['titulo']}: estudio de feria"
    doc.core_properties.author = 'Alliance Albors y Prado S.L.'
    doc.save(ruta)


# ----------------------------------------------------------------- Excel y CSV

def hoja(wb, nombre, cabecera, filas, anchos):
    ws = wb.create_sheet(nombre)
    ws.append(cabecera)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor=GRIS_CABECERA)
        c.alignment = Alignment(vertical='top', wrap_text=True)
    for f in filas:
        ws.append(f)
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def filas_contactos(d):
    filas = []
    for e in sorted(d['empresas'], key=lambda e: e['prioridad']):
        for c in sorted(e['contactos'], key=lambda c: c['perfil']):
            filas.append([e['prioridad'], e['nombre'], e['categoria'], c['nombre'], c['cargo'], c['perfil'],
                          PERFILES[c['perfil']], c['ubicacion'], c['origen'], c['linkedin'], c['certeza'],
                          e['telefono'], e['email'], e['web'], '', '', ''])
    return filas


def dividir_nombre(nombre):
    partes = nombre.split()
    return partes[0], ' '.join(partes[1:])


def filas_apollo(d):
    filas = []
    for e in sorted(d['empresas'], key=lambda e: e['prioridad']):
        for c in sorted(e['contactos'], key=lambda c: c['perfil']):
            nombre, apellidos = dividir_nombre(c['nombre'])
            empresa = e['nombre'].split(' (')[0]
            linkedin = f"https://www.{c['linkedin']}" if c['linkedin'] and not c['linkedin'].startswith('http') \
                else c['linkedin']
            filas.append([nombre, apellidos, c['cargo'], empresa, e['web'], linkedin, c['ubicacion'],
                          f"P{e['prioridad']}-{c['perfil']}"])
    return filas


def generar_excel(d, ruta_xlsx, ruta_csv):
    wb = Workbook()
    wb.remove(wb.active)
    hoja(wb, 'Contactos',
         ['Prioridad', 'Empresa', 'Categoría', 'Nombre', 'Cargo', 'Perfil', 'Descripción del perfil', 'Ubicación',
          'Origen del dato', 'LinkedIn', 'Certeza', 'Teléfono empresa', 'Correo empresa', 'Web',
          'Correo personal', 'Teléfono directo', 'Notas'],
         filas_contactos(d), [8, 26, 22, 24, 40, 7, 30, 12, 30, 34, 12, 26, 24, 22, 26, 18, 30])
    hoja(wb, 'Empresas',
         ['Prioridad', 'Empresa', 'Categoría', 'Papel', 'Centros de datos', 'Agua, PCI y refrigeración', 'Clientes',
          'Ferias y foros', 'Oficinas', 'Teléfono', 'Correo general', 'Web', 'Por qué'],
         [[e['prioridad'], e['nombre'], e['categoria'], e['rol'], e['dc'], e['agua_pci'], e['clientes'], e['feria'],
           e['oficinas'], e['telefono'], e['email'], e['web'], e['motivo']] for e in d['empresas']],
         [8, 26, 22, 50, 26, 36, 26, 36, 32, 22, 22, 20, 36])
    hoja(wb, 'Proyectos',
         ['Promotor', 'Ubicación', 'Potencia', 'Inversión', 'Estado', 'Ingeniería y construcción', 'Interés para Stora'],
         [[p['promotor'], p['ubicacion'], p['potencia'], p['inversion'], p['estado'], p['agentes'], p['interes']]
          for p in d['mercado']['proyectos']],
         [26, 34, 22, 22, 34, 50, 34])
    hoja(wb, 'Operadores', ['Operador', 'Nombre', 'Cargo', 'Dónde consta'], d['operadores'], [22, 24, 40, 26])
    cab_apollo = ['First Name', 'Last Name', 'Title', 'Company', 'Website', 'Person Linkedin Url', 'City', 'Prioridad Stora']
    apollo = filas_apollo(d)
    hoja(wb, 'Para Apollo', cab_apollo, apollo, [16, 24, 44, 30, 24, 44, 14, 14])
    wb.save(ruta_xlsx)
    with open(ruta_csv, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow(cab_apollo)
        w.writerows(apollo)


# ----------------------------------------------------------------- principal

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--fecha', default=datetime.date.today().isoformat(),
                    help='fecha del nombre de archivo (AAAA-MM-DD); por defecto, hoy')
    args = ap.parse_args()
    try:
        fecha = datetime.date.fromisoformat(args.fecha)
    except ValueError:
        ap.error('la fecha debe tener el formato AAAA-MM-DD')

    with open(DATOS_ESTUDIO, encoding='utf-8') as f:
        d = json.load(f)
    errores, avisos = comprobar(d)
    for m in avisos:
        print('AVISO  ', m)
    if errores:
        for m in errores:
            print('ERROR  ', m)
        print(f'\nComprobación fallida: {len(errores)} errores. No se genera ningún documento.')
        return 1

    n = contar(d)
    est = d['estudio']
    base = f"{est['archivo']}_{fecha.isoformat()}"
    docx = SALIDAS / f'{base}.docx'
    SALIDAS.mkdir(exist_ok=True)
    generar_docx(d, docx, fecha.strftime('%d/%m/%Y'), n)
    pdf = bloc_pdf.convertir_a_pdf(docx, SALIDAS)
    base_c = f"{est['archivo_contactos']}_{fecha.isoformat()}"
    xlsx, csv_apollo = SALIDAS / f'{base_c}.xlsx', SALIDAS / f'{base_c}_Apollo.csv'
    generar_excel(d, xlsx, csv_apollo)
    for ruta in (docx, pdf, xlsx, csv_apollo):
        print(f'Generado: {Path(ruta).relative_to(RAIZ)}')
    print(f"{n['n_empresas']} empresas · {n['n_contactos']} contactos · {len(d['mercado']['proyectos'])} proyectos.")
    return 0


if __name__ == '__main__':
    sys.exit(main())
