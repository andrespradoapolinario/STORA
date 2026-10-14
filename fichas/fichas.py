"""Fichas maestras: un expediente por empresa y una ficha en Word y PDF.

Uso:
    python fichas.py inicializar --organizaciones ORG.json --salida expedientes.json
    python fichas.py añadir expedientes.json entradas.json [--simular] [--organizaciones ORG.json]
    python fichas.py generar expedientes.json SALIDA --organizaciones ORG.json
                             --panel PANEL.xlsx --tracker TRACKER.xlsx [--todas]
    python fichas.py validar expedientes.json --organizaciones ORG.json
                             [--panel PANEL.xlsx] [--tracker TRACKER.xlsx] [--salida SALIDA]

Formato de entradas.json, procedimiento y conflictos: INSTRUCCIONES_FICHAS.md.

El expediente solo guarda lo que no está en otra fuente: el perfil y los
contactos se leen de organizaciones.json, el estado y el próximo compromiso del
panel de seguimiento y la valoración del tracker de la feria. El panel y el
tracker solo se leen; nunca se escriben.

El módulo es autónomo (no importa nada de scripts/) para que funcione copiado
fuera del repositorio. Necesita Python 3.10 o posterior, openpyxl y
LibreOffice; con pdfplumber comprueba además la paginación de cada PDF.
"""
import argparse
import copy
import datetime
import hashlib
import importlib.util
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape, quoteattr

AQUI = Path(__file__).resolve().parent
ESQUEMA = AQUI / 'esquema_expediente.json'
PLANTILLA = AQUI / 'plantilla'
VERSION = 1
ESTADO = 'estado_fichas.json'
CARPETA_FICHAS = 'fichas'
INDICE = 'Indice_Fichas_Maestras'
COPIAS = 'copias_seguridad'

GRUPOS = {
    'A': 'Promotoras',
    'B': 'EPC',
    'C': 'Ingenierías',
    'D': 'Consultorías',
    'E': 'Cooperativas y asociaciones',
}
SUBTIPOS_ORGANIZACION = ('Asociación', 'Organismo público', 'Centro de investigación')

# Listas del expediente y prefijo de los id que les asigna el script.
LISTAS = {
    'lo_que_sabemos': 'sab', 'proyectos': 'pro', 'tanques': 'tan', 'cotizaciones': 'cot',
    'competencia': 'com', 'pendientes': 'pen', 'cronologia': 'cro', 'documentos': 'doc',
}
CAMPOS_TANQUE = ('proyecto', 'uso', 'cantidad', 'diametro_m', 'altura_m', 'volumen_m3', 'techo',
                 'aislamiento', 'otros')
# Duplicados: dos entradas con los mismos campos de clave son la misma, y la
# nueva se omite. Si además difieren en un campo de contenido, la nueva es un
# intento de modificar la existente (conflicto). La cronología sigue su propia
# regla (misma referencia de Gmail, o misma fecha, tipo y resumen) y una entrada
# repetida se omite siempre.
CLAVES = {
    'lo_que_sabemos': (('texto',), ('naturaleza',)),
    'proyectos': (('nombre', 'situacion'), ('ubicacion',)),
    'tanques': (CAMPOS_TANQUE, ()),
    'cotizaciones': (('fecha', 'descripcion'), ('estado', 'importe_eur')),
    'competencia': (('texto',), ()),
    'pendientes': (('accion', 'de_quien'), ('fecha_limite',)),
    'documentos': (('nombre', 'fecha'), ('ruta',)),
}

TIPOS = {
    'correo_recibido': 'Correo recibido', 'correo_enviado': 'Correo enviado', 'reunion': 'Reunión',
    'llamada': 'Llamada', 'nota': 'Nota', 'whatsapp': 'WhatsApp', 'documento': 'Documento',
}
FUENTES = {'gmail': 'Gmail', 'conversacion': 'Conversación', 'foto_nota': 'Foto de nota', 'documento': 'Documento'}
NATURALEZAS = {'dato': 'DATO', 'estimacion': 'ESTIMACIÓN', 'supuesto': 'SUPUESTO'}
DE_QUIEN = {'nosotros': 'Nosotros', 'Stora': 'Stora', 'cliente': 'Cliente'}

# Panel de seguimiento y tracker de la feria: hojas y columnas (la fila de
# cabecera se localiza por su texto).
HOJAS_PANEL = ('PROMOTORES', 'EPC', 'INGENIERÍAS Y CONSULTORÍAS', 'OTROS')
COLUMNAS_PANEL = {'nombre': 'Empresa', 'estado': 'Estado', 'fecha': 'Fecha', 'contacto': 'Contacto',
                  'telefono': 'Teléfono', 'email': 'Email', 'notas': 'Notas'}
HOJAS_TRACKER = ('Lead Overview',)
COLUMNAS_TRACKER = {'rating': 'Rating', 'nombre': 'Company', 'evento': 'Exhibition / Event',
                    'contacto': 'Contact', 'oportunidad': 'Opportunity', 'fase': 'Stage'}
FILAS_BUSQUEDA_CABECERA = 30
# Valoraciones en texto que se reconocen para ordenar (mayor = mejor). Los
# números y las estrellas (★) se ordenan por su valor.
PALABRAS_VALORACION = {
    'hot': 3, 'caliente': 3, 'alta': 3, 'alto': 3, 'high': 3, 'a': 3,
    'warm': 2, 'templado': 2, 'templada': 2, 'media': 2, 'medio': 2, 'medium': 2, 'b': 2,
    'cold': 1, 'frio': 1, 'fria': 1, 'baja': 1, 'bajo': 1, 'low': 1, 'c': 1,
}

FECHA = re.compile(r'^\d{4}-\d{2}-\d{2}$')
ID_EMPRESA = re.compile(r'^[a-z0-9]+(-[a-z0-9]+)*$')
FORMAS_JURIDICAS = re.compile(r'( (s l u|s l l|s l|s a u|s a|slu|sll|sl|sau|sa))+$')
CONTROL = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f]')
LARGO_SITUACION = 700  # caracteres; la situación actual debe ocupar 3-5 líneas


class ErrorDatos(Exception):
    pass


class ErrorGeneracion(Exception):
    pass


# --- Utilidades ----------------------------------------------------------------

def sin_tildes(texto):
    return unicodedata.normalize('NFKD', texto).encode('ascii', 'ignore').decode()


def normalizar(texto):
    """Minúsculas, sin tildes ni puntuación y con los espacios simplificados."""
    return ' '.join(re.findall(r'[a-z0-9]+', sin_tildes(str(texto)).casefold()))


def variantes_nombre(nombre):
    """Formas con las que se compara el nombre de una empresa: sin mayúsculas,
    tildes ni forma jurídica (S.L., S.A.), con el texto entre paréntesis
    suprimido y conservado."""
    variantes = set()
    for base in (re.sub(r'\([^)]*\)', ' ', nombre), nombre.replace('(', ' ').replace(')', ' ')):
        v = FORMAS_JURIDICAS.sub('', normalizar(base))
        if v:
            variantes.add(v)
    return variantes


def crear_id(nombre):
    """«Biorig (grupo Solarig)» -> «biorig» (misma regla que scripts/comun.py)."""
    base = re.sub(r'\s*\([^)]*\)', '', nombre).replace('Ø', '0').replace('ø', '0')
    base = unicodedata.normalize('NFKD', base).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '-', base.lower()).strip('-')


def vacio(valor):
    return valor is None or valor == '' or valor == []


def fecha_valida(texto):
    if not isinstance(texto, str) or not FECHA.match(texto):
        return False
    try:
        datetime.date.fromisoformat(texto)
    except ValueError:
        return False
    return True


def fecha_es(valor):
    """«2026-10-14» -> «14/10/2026». Un texto que no es fecha ISO se devuelve tal cual."""
    if isinstance(valor, datetime.datetime):
        valor = valor.date()
    if isinstance(valor, datetime.date):
        return valor.strftime('%d/%m/%Y')
    if fecha_valida(valor):
        return datetime.date.fromisoformat(valor).strftime('%d/%m/%Y')
    return '' if valor is None else str(valor).strip()


def numero_es(valor):
    """6000 -> «6.000»; 18.5 -> «18,5»."""
    if vacio(valor):
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = f'{valor:,}' if isinstance(valor, int) else f'{valor:,.2f}'.rstrip('0').rstrip('.')
    return texto.replace(',', ' ').replace('.', ',').replace(' ', '.')


def leer_json(ruta):
    with open(ruta, encoding='utf-8') as f:
        return json.load(f)


def escribir_json(datos, ruta):
    ruta = Path(ruta)
    tmp = ruta.with_name(ruta.name + '.tmp')
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(datos, f, ensure_ascii=False, indent=2)
        f.write('\n')
    tmp.replace(ruta)


def copia_seguridad(ruta):
    """Copia fechada del archivo en copias_seguridad/, junto a él."""
    ruta = Path(ruta)
    carpeta = ruta.parent / COPIAS
    carpeta.mkdir(exist_ok=True)
    sello = datetime.datetime.now().strftime('%Y-%m-%d_%H%M%S')
    destino = carpeta / f'{ruta.stem}_{sello}{ruta.suffix}'
    k = 2
    while destino.exists():
        destino = carpeta / f'{ruta.stem}_{sello}_{k}{ruta.suffix}'
        k += 1
    shutil.copy2(ruta, destino)
    return destino


def cargar_organizaciones(ruta):
    try:
        datos = leer_json(ruta)
    except (OSError, json.JSONDecodeError) as exc:
        raise ErrorDatos(f'No se puede leer {Path(ruta).name}: {exc}') from exc
    orgs = datos.get('organizaciones') if isinstance(datos, dict) else datos
    if not isinstance(orgs, list) or not all(isinstance(o, dict) and o.get('id') for o in orgs):
        raise ErrorDatos(f'{Path(ruta).name} no tiene una lista «organizaciones» con «id».')
    return orgs


def fecha_organizaciones(ruta, indicada=None):
    """Fecha de la copia de organizaciones.json: la indicada, la del nombre del
    archivo (organizaciones_AAAA-MM-DD.json) o la de su última modificación."""
    if indicada:
        return indicada
    m = re.search(r'\d{4}-\d{2}-\d{2}', Path(ruta).name)
    if m and fecha_valida(m.group(0)):
        return m.group(0)
    return datetime.date.fromtimestamp(Path(ruta).stat().st_mtime).isoformat()


# --- Esquema -------------------------------------------------------------------

NOMBRES_TIPO = {'object': 'un objeto', 'array': 'una lista', 'string': 'un texto', 'integer': 'un número entero',
                'number': 'un número', 'null': 'null', 'boolean': 'true o false'}
_esquema = None


def esquema():
    global _esquema
    if _esquema is None:
        _esquema = leer_json(ESQUEMA)
    return _esquema


def _es_tipo(valor, tipo):
    return {
        'object': isinstance(valor, dict),
        'array': isinstance(valor, list),
        'string': isinstance(valor, str),
        'integer': isinstance(valor, int) and not isinstance(valor, bool),
        'number': isinstance(valor, (int, float)) and not isinstance(valor, bool),
        'null': valor is None,
        'boolean': isinstance(valor, bool),
    }[tipo]


def _resolver(sub):
    while '$ref' in sub:
        sub = esquema()['$defs'][sub['$ref'].removeprefix('#/$defs/')]
    return sub


def _tipo_admitido(valor, sub):
    sub = _resolver(sub)
    tipos = sub.get('type')
    if tipos is None:
        return True
    return any(_es_tipo(valor, t) for t in ([tipos] if isinstance(tipos, str) else tipos))


def errores_esquema(valor, sub=None, ruta='expedientes'):
    """Valida con el subconjunto de JSON Schema que usa esquema_expediente.json.

    Devuelve la lista de errores, con la ruta de cada uno.
    """
    sub = _resolver(esquema() if sub is None else sub)
    errores = []
    for parte in sub.get('allOf', []):
        errores += errores_esquema(valor, parte, ruta)
    if 'anyOf' in sub:
        opciones = sub['anyOf']
        if all(errores_esquema(valor, o, ruta) for o in opciones):
            elegida = next((o for o in opciones if _tipo_admitido(valor, o)), opciones[-1])
            errores += errores_esquema(valor, elegida, ruta)
    tipos = sub.get('type')
    if tipos:
        tipos = [tipos] if isinstance(tipos, str) else tipos
        if not any(_es_tipo(valor, t) for t in tipos):
            return errores + [f'{ruta}: debe ser {" o ".join(NOMBRES_TIPO[t] for t in tipos)}.']
    if 'const' in sub and valor != sub['const']:
        errores.append(f'{ruta}: debe valer {json.dumps(sub["const"])}.')
    if 'enum' in sub and valor not in sub['enum']:
        errores.append(f'{ruta}: «{valor}» no es un valor admitido ({", ".join(map(str, sub["enum"]))}).')
    if isinstance(valor, str):
        if len(valor.strip()) < sub.get('minLength', 0):
            errores.append(f'{ruta}: no puede estar vacío.')
        elif 'pattern' in sub and not re.search(sub['pattern'], valor):
            errores.append(f'{ruta}: «{valor}» no tiene el formato esperado.')
        elif sub.get('format') == 'date' and not fecha_valida(valor):
            errores.append(f'{ruta}: «{valor}» no es una fecha válida (AAAA-MM-DD).')
    if _es_tipo(valor, 'number'):
        if 'minimum' in sub and valor < sub['minimum']:
            errores.append(f'{ruta}: debe ser como mínimo {sub["minimum"]}.')
        if 'exclusiveMinimum' in sub and valor <= sub['exclusiveMinimum']:
            errores.append(f'{ruta}: debe ser mayor que {sub["exclusiveMinimum"]}.')
    if isinstance(valor, dict):
        for campo in sub.get('required', []):
            if campo not in valor:
                errores.append(f'{ruta}: falta «{campo}».')
        propiedades = sub.get('properties', {})
        adicionales = sub.get('additionalProperties', True)
        for campo, v in valor.items():
            if campo in propiedades:
                errores += errores_esquema(v, propiedades[campo], f'{ruta}.{campo}')
            elif adicionales is False:
                errores.append(f'{ruta}: campo desconocido «{campo}».')
            elif isinstance(adicionales, dict):
                errores += errores_esquema(v, adicionales, f'{ruta}.{campo}')
    if isinstance(valor, list) and 'items' in sub:
        for i, v in enumerate(valor):
            errores += errores_esquema(v, sub['items'], f'{ruta}[{i}]')
    return errores


def incoherencias(lista, entrada):
    """Reglas de una entrada que el esquema no expresa."""
    m = []
    if lista == 'pendientes':
        if entrada.get('estado') == 'cerrado' and not entrada.get('fecha_cierre'):
            m.append('un pendiente cerrado necesita «fecha_cierre».')
        if entrada.get('estado') == 'abierto' and entrada.get('fecha_cierre'):
            m.append('un pendiente abierto no lleva «fecha_cierre» (para cerrarlo, use «cerrar_pendientes»).')
    if lista == 'cronologia' and entrada.get('fuente', {}).get('tipo') == 'gmail' \
            and not entrada['fuente'].get('ref', '').strip():
        m.append('un correo de Gmail necesita «fuente.ref» (identificador del mensaje).')
    if lista == 'tanques' and all(vacio(entrada.get(c)) for c in CAMPOS_TANQUE):
        m.append('el tanque no tiene ningún dato técnico.')
    return m


def errores_expedientes(exp):
    """Esquema y reglas adicionales: clave = id, id de entradas únicos y coherencia."""
    errores = errores_esquema(exp)
    if errores:
        return errores
    for clave, e in exp['empresas'].items():
        if not ID_EMPRESA.match(clave):
            errores.append(f'empresas.{clave}: clave no válida (minúsculas, números y guiones).')
        if e['id'] != clave:
            errores.append(f'empresas.{clave}: el id «{e["id"]}» no coincide con la clave.')
        for lista in LISTAS:
            vistos = set()
            for entrada in e[lista]:
                if entrada['id'] in vistos:
                    errores.append(f'empresas.{clave}.{lista}: id repetido «{entrada["id"]}».')
                vistos.add(entrada['id'])
                errores += [f'empresas.{clave}.{lista} {entrada["id"]}: {m}' for m in incoherencias(lista, entrada)]
    return errores


def cargar_expedientes(ruta):
    """(expedientes, errores)."""
    try:
        exp = leer_json(ruta)
    except OSError as exc:
        return None, [f'No se puede leer {Path(ruta).name}: {exc}']
    except json.JSONDecodeError as exc:
        return None, [f'{Path(ruta).name}: JSON no válido ({exc}).']
    return exp, errores_expedientes(exp)


def expediente_vacio(oid):
    return {'id': oid, 'nombres_alternativos': [], 'situacion_actual': None,
            **{lista: [] for lista in LISTAS}}


# --- Asociación de nombres -----------------------------------------------------

class Asociador:
    """Asocia un nombre (del panel, del tracker o de las entradas) con el id de una
    empresa: primero por id, después por el nombre de organizaciones.json y por
    último por los nombres alternativos del expediente."""

    def __init__(self, orgs, empresas):
        self.ids = {o['id'] for o in orgs} | set(empresas)
        self.nombres, self.alternativos = {}, {}
        for o in orgs:
            for v in variantes_nombre(o.get('nombre', '')):
                self.nombres.setdefault(v, set()).add(o['id'])
        for oid, e in empresas.items():
            for nombre in e.get('nombres_alternativos', []):
                self.anadir_alternativo(oid, nombre)

    def anadir_alternativo(self, oid, nombre):
        for v in variantes_nombre(nombre):
            self.alternativos.setdefault(v, set()).add(oid)

    def buscar(self, nombre):
        """(id, '') si el nombre corresponde a una sola empresa; (None, motivo) si no."""
        texto = str(nombre).strip()
        for candidato in (texto, texto.lower(), crear_id(texto)):
            if candidato in self.ids:
                return candidato, ''
        variantes = variantes_nombre(texto)
        for indice in (self.nombres, self.alternativos):
            ids = set().union(*(indice.get(v, set()) for v in variantes))
            if len(ids) == 1:
                return ids.pop(), ''
            if len(ids) > 1:
                return None, f'«{texto}» corresponde a varias empresas ({", ".join(sorted(ids))}).'
        return None, (f'«{texto}» no se puede asociar a ninguna empresa '
                      '(ni por id, ni por nombre, ni por nombre alternativo).')


# --- inicializar ---------------------------------------------------------------

def inicializar(ruta_org, ruta_exp, informe=print):
    ruta_exp = Path(ruta_exp)
    try:
        orgs = cargar_organizaciones(ruta_org)
    except ErrorDatos as exc:
        informe(f'ERROR      {exc}')
        return 1
    if ruta_exp.exists():
        exp, errores = cargar_expedientes(ruta_exp)
        if errores:
            for m in errores:
                informe(f'ERROR      {m}')
            informe(f'\n{ruta_exp.name} no es válido: no se ha modificado.')
            return 1
    else:
        exp = {'version': VERSION, 'empresas': {}}
    nuevas = [o['id'] for o in orgs if o['id'] not in exp['empresas']]
    for oid in nuevas:
        exp['empresas'][oid] = expediente_vacio(oid)
    ids = {o['id'] for o in orgs}
    for oid in exp['empresas']:
        if oid not in ids:
            informe(f'AVISO      expediente «{oid}» sin correspondencia en {Path(ruta_org).name}.')
    if nuevas:
        if ruta_exp.exists():
            informe(f'Copia de seguridad: {copia_seguridad(ruta_exp)}')
        escribir_json(exp, ruta_exp)
    informe(f'{len(orgs)} organizaciones en {Path(ruta_org).name} · {len(nuevas)} expedientes creados · '
            f'{len(orgs) - len(nuevas)} ya existían · {len(exp["empresas"])} en {ruta_exp.name}.')
    return 0


# --- añadir --------------------------------------------------------------------

def comparable(valor):
    if isinstance(valor, str):
        return normalizar(valor)
    if isinstance(valor, list):
        return tuple(comparable(v) for v in valor)
    if isinstance(valor, dict):
        return tuple(sorted((k, comparable(v)) for k, v in valor.items()))
    if isinstance(valor, bool) or valor is None:
        return '' if vacio(valor) else valor
    return float(valor)


def diferencias(a, b, campos):
    return [c for c in campos if comparable(a.get(c)) != comparable(b.get(c))]


def gmail_ref(entrada):
    fuente = entrada.get('fuente') or {}
    return fuente.get('ref', '').strip() if fuente.get('tipo') == 'gmail' else ''


def buscar_duplicada(existentes, lista, entrada):
    """(entrada existente, motivo) si la nueva ya consta; (None, '') si no."""
    if lista == 'cronologia':
        ref = gmail_ref(entrada)
        clave = (entrada['fecha'], entrada['tipo'], normalizar(entrada['resumen']))
        for x in existentes:
            if ref and gmail_ref(x) == ref:
                return x, f'misma referencia de Gmail {ref}'
            if (x['fecha'], x['tipo'], normalizar(x['resumen'])) == clave:
                return x, 'misma fecha, tipo y resumen'
        return None, ''
    campos = CLAVES[lista][0]
    clave = tuple(comparable(entrada.get(c)) for c in campos)
    for x in existentes:
        if tuple(comparable(x.get(c)) for c in campos) == clave:
            return x, f'coincide en {", ".join(campos)}'
    return None, ''


def siguiente_id(existentes, prefijo):
    n = max((int(x['id'].split('-')[1]) for x in existentes), default=0)
    return f'{prefijo}-{n + 1:04d}'


def describir(lista, x):
    if lista == 'cronologia':
        texto = f'{fecha_es(x["fecha"])} · {TIPOS[x["tipo"]]}: {x["resumen"]}'
    else:
        campo = next(c for c in ('texto', 'nombre', 'descripcion', 'accion', 'proyecto', 'uso', 'fuente')
                     if x.get(c))
        texto = str(x[campo])
    return texto if len(texto) <= 80 else texto[:77] + '…'


def anadir_entrada(e, lista, entrada, ref, informe):
    desc = f'{ref} · {lista}'
    if 'id' in entrada:
        existente = next((x for x in e[lista] if x['id'] == entrada['id']), None)
        if existente is None:
            informe.append(('ERROR', f'{desc}: el id {entrada["id"]} no existe. Los id los asigna el script: '
                                     'las entradas nuevas van sin «id».'))
        else:
            campos = sorted((set(existente) | set(entrada)) - {'id'})
            cambios = diferencias(existente, entrada, campos)
            if cambios:
                informe.append(('CONFLICTO', f'{desc}: intento de modificar la entrada {entrada["id"]} '
                                             f'(cambia {", ".join(cambios)}). '
                                             'Las entradas existentes no se modifican.'))
            else:
                informe.append(('omitida', f'{desc}: {entrada["id"]} ya consta con el mismo contenido.'))
        return
    problemas = incoherencias(lista, entrada)
    if problemas:
        informe += [('ERROR', f'{desc} «{describir(lista, entrada)}»: {m}') for m in problemas]
        return
    existente, motivo = buscar_duplicada(e[lista], lista, entrada)
    if existente:
        if lista == 'pendientes' and entrada['estado'] == 'cerrado' and existente['estado'] == 'abierto':
            informe.append(('CONFLICTO', f'{desc}: {existente["id"]} «{describir(lista, existente)}» sigue abierto. '
                                         'Para cerrarlo, use «cerrar_pendientes».'))
            return
        cambios = diferencias(existente, entrada, CLAVES.get(lista, ((), ()))[1])
        if cambios:
            informe.append(('CONFLICTO', f'{desc}: intento de modificar la entrada {existente["id"]} '
                                         f'«{describir(lista, existente)}» (cambia {", ".join(cambios)}). '
                                         'Las entradas existentes no se modifican.'))
        else:
            informe.append(('omitida', f'{desc}: duplicada de {existente["id"]} ({motivo}).'))
        return
    nueva = {'id': siguiente_id(e[lista], LISTAS[lista]), **entrada}
    e[lista].append(nueva)
    informe.append(('añadida', f'{desc}: {nueva["id"]} «{describir(lista, nueva)}».'))


def cerrar_pendiente(e, cierre, ref, informe):
    desc = f'{ref} · cerrar_pendientes'
    if cierre.get('id'):
        candidatos = [p for p in e['pendientes'] if p['id'] == cierre['id']]
        quien = cierre['id']
    elif cierre.get('accion'):
        candidatos = [p for p in e['pendientes'] if normalizar(p['accion']) == normalizar(cierre['accion'])]
        quien = f'«{cierre["accion"]}»'
        if len(candidatos) > 1:
            candidatos = [p for p in candidatos if p['estado'] == 'abierto'] or candidatos
    else:
        informe.append(('ERROR', f'{desc}: indique el «id» o la «accion» del pendiente.'))
        return
    if not candidatos:
        informe.append(('ERROR', f'{desc}: no hay ningún pendiente {quien}.'))
        return
    if len(candidatos) > 1:
        informe.append(('ERROR', f'{desc}: {quien} corresponde a varios pendientes '
                                 f'({", ".join(p["id"] for p in candidatos)}); indique el id.'))
        return
    p = candidatos[0]
    if p['estado'] == 'cerrado':
        if p.get('fecha_cierre') == cierre['fecha_cierre']:
            informe.append(('omitida', f'{desc}: {p["id"]} ya estaba cerrado el {fecha_es(p["fecha_cierre"])}.'))
        else:
            informe.append(('CONFLICTO', f'{desc}: {p["id"]} ya está cerrado desde el {fecha_es(p["fecha_cierre"])}; '
                                         f'las entradas proponen el {fecha_es(cierre["fecha_cierre"])}.'))
        return
    p['estado'] = 'cerrado'
    p['fecha_cierre'] = cierre['fecha_cierre']
    informe.append(('cambio', f'{desc}: {p["id"]} «{describir("pendientes", p)}» cerrado el '
                              f'{fecha_es(cierre["fecha_cierre"])}.'))


def aplicar(exp, bloques, orgs=()):
    """Aplica las entradas sobre una copia del expediente.

    Devuelve (expediente_nuevo, informe). El informe es la lista ordenada de
    (tipo, texto), con tipo «añadida», «cambio», «omitida», «AVISO»,
    «CONFLICTO» o «ERROR». Con algún conflicto o error no debe guardarse nada.
    """
    nuevo = copy.deepcopy(exp)
    informe = []
    asociador = Asociador(orgs, nuevo['empresas'])
    nombres = {o['id']: o.get('nombre', o['id']) for o in orgs}
    for n, bloque in enumerate(bloques, 1):
        oid, motivo = asociador.buscar(bloque['empresa'])
        if oid is None:
            informe.append(('CONFLICTO', f'bloque {n}: {motivo}'))
            continue
        if oid not in nuevo['empresas']:
            informe.append(('CONFLICTO', f'bloque {n}: «{bloque["empresa"]}» ({oid}) no tiene expediente. '
                                         'Ejecute «inicializar» con el organizaciones.json actual.'))
            continue
        e = nuevo['empresas'][oid]
        ref = f'{nombres[oid]} [{oid}]' if oid in nombres else oid
        for nombre in bloque.get('nombres_alternativos', []):
            otro, _ = asociador.buscar(nombre)
            if otro == oid:
                informe.append(('omitida', f'{ref}: el nombre «{nombre}» ya identifica a esta empresa.'))
            elif otro:
                informe.append(('CONFLICTO', f'{ref}: el nombre alternativo «{nombre}» ya identifica a «{otro}».'))
            else:
                e['nombres_alternativos'].append(nombre.strip())
                asociador.anadir_alternativo(oid, nombre)
                informe.append(('cambio', f'{ref}: nuevo nombre alternativo «{nombre.strip()}».'))
        situacion = bloque.get('situacion_actual')
        if situacion:
            actual = e['situacion_actual']
            if actual and normalizar(actual['texto']) == normalizar(situacion['texto']):
                informe.append(('omitida', f'{ref}: la situación actual ya constaba.'))
            elif actual and actual['fecha'] > situacion['fecha']:
                informe.append(('CONFLICTO', f'{ref}: la situación actual registrada ({fecha_es(actual["fecha"])}) '
                                             f'es posterior a la propuesta ({fecha_es(situacion["fecha"])}).'))
            else:
                e['situacion_actual'] = {'texto': situacion['texto'].strip(), 'fecha': situacion['fecha']}
                informe.append(('cambio', f'{ref}: situación actual {"reescrita" if actual else "nueva"} '
                                          f'({fecha_es(situacion["fecha"])}).'))
                if len(situacion['texto']) > LARGO_SITUACION:
                    informe.append(('AVISO', f'{ref}: la situación actual tiene {len(situacion["texto"])} caracteres; '
                                             'debe ocupar de 3 a 5 líneas.'))
        for lista in LISTAS:
            for entrada in bloque.get(lista, []):
                anadir_entrada(e, lista, entrada, ref, informe)
        for cierre in bloque.get('cerrar_pendientes', []):
            cerrar_pendiente(e, cierre, ref, informe)
    return nuevo, informe


def anadir(ruta_exp, ruta_entradas, simular=False, ruta_org=None, informe=print):
    ruta_exp = Path(ruta_exp)
    exp, errores = cargar_expedientes(ruta_exp)
    if errores:
        for m in errores:
            informe(f'ERROR      {m}')
        informe(f'\n{ruta_exp.name} no es válido: no se ha aplicado nada.')
        return 1
    try:
        bloques = leer_json(ruta_entradas)
        orgs = cargar_organizaciones(ruta_org) if ruta_org else []
    except (OSError, json.JSONDecodeError, ErrorDatos) as exc:
        informe(f'ERROR      {exc}')
        informe('\nNo se ha aplicado nada.')
        return 1
    errores = errores_esquema(bloques, {'$ref': '#/$defs/entradas'}, 'entradas')
    if errores:
        for m in errores:
            informe(f'ERROR      {m}')
        informe(f'\n{len(errores)} error(es) de formato en {Path(ruta_entradas).name}: no se ha aplicado nada.')
        return 1
    nuevo, incidencias = aplicar(exp, bloques, orgs)
    if not any(t in ('CONFLICTO', 'ERROR') for t, _ in incidencias):
        incidencias += [('ERROR', f'resultado: {m}') for m in errores_expedientes(nuevo)]
    informe(f'== {Path(ruta_entradas).name} → {ruta_exp.name}')
    for tipo, texto in incidencias:
        if tipo in ('añadida', 'cambio'):
            informe(f'  · {texto}')
        else:
            informe(f'{tipo.upper():<10} {texto}')
    cuenta = {t: sum(1 for x, _ in incidencias if x == t)
              for t in ('añadida', 'cambio', 'omitida', 'CONFLICTO', 'ERROR')}
    informe(f'\n{cuenta["añadida"]} entradas añadidas · {cuenta["cambio"]} cambios · {cuenta["omitida"]} omitidas '
            f'· {cuenta["CONFLICTO"]} conflicto(s) · {cuenta["ERROR"]} error(es).')
    if cuenta['CONFLICTO'] or cuenta['ERROR']:
        informe('No se ha aplicado nada: corrija los errores y consulte a Andrés los conflictos antes de repetir.')
        return 1 if cuenta['ERROR'] else 2
    if not cuenta['añadida'] and not cuenta['cambio']:
        informe('Sin cambios: todo lo que traen las entradas ya constaba.')
        return 0
    if simular:
        informe('Simulación: no se ha modificado ningún archivo.')
        return 0
    informe(f'Copia de seguridad: {copia_seguridad(ruta_exp)}')
    escribir_json(nuevo, ruta_exp)
    informe(f'Expediente actualizado: {ruta_exp}')
    return 0


# --- Panel y tracker -----------------------------------------------------------

def _posiciones_cabecera(valores, columnas, obligatorias):
    """{campo: índice de columna} si la fila es la cabecera; None si no."""
    textos = [normalizar(v) if isinstance(v, str) else '' for v in valores]
    posiciones = {}
    for campo, titulo in columnas.items():
        t = normalizar(titulo)
        exacta = [i for i, x in enumerate(textos) if x == t]
        prefijo = [i for i, x in enumerate(textos) if x.startswith(t + ' ')]
        if exacta or prefijo:
            posiciones[campo] = (exacta or prefijo)[0]
    return posiciones if all(c in posiciones for c in obligatorias) else None


def leer_excel(ruta, hojas, columnas, obligatorias):
    """Filas con nombre de las hojas indicadas: [{campo: valor, 'hoja', 'fila'}].

    La fila de cabecera se localiza por su texto. El libro se abre solo para lectura.
    """
    import openpyxl

    ruta = Path(ruta)
    try:
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza tipos variados con un archivo dañado
        raise ErrorDatos(f'No se puede leer {ruta.name}: {exc}') from exc
    filas, avisos = [], []
    try:
        por_nombre = {normalizar(h): h for h in libro.sheetnames}
        encontradas = [por_nombre[normalizar(h)] for h in hojas if normalizar(h) in por_nombre]
        for h in hojas:
            if normalizar(h) not in por_nombre:
                avisos.append(f'{ruta.name}: no hay hoja «{h}».')
        if not encontradas:
            raise ErrorDatos(f'{ruta.name}: no tiene ninguna de las hojas {", ".join(hojas)}.')
        for nombre_hoja in encontradas:
            posiciones = None
            for n, valores in enumerate(libro[nombre_hoja].iter_rows(values_only=True), 1):
                if posiciones is None:
                    if n > FILAS_BUSQUEDA_CABECERA:
                        break
                    posiciones = _posiciones_cabecera(valores, columnas, obligatorias)
                    continue
                fila = {c: (valores[i] if i < len(valores) else None) for c, i in posiciones.items()}
                if isinstance(fila['nombre'], str):
                    fila['nombre'] = fila['nombre'].strip()
                if vacio(fila['nombre']):
                    continue
                filas.append({**fila, 'hoja': nombre_hoja, 'fila': n})
            if posiciones is None:
                raise ErrorDatos(f'{ruta.name}, hoja «{nombre_hoja}»: no se encuentra la fila de cabecera '
                                 f'(columnas {", ".join(columnas[c] for c in obligatorias)}).')
    finally:
        libro.close()
    return filas, avisos


def leer_panel(ruta):
    return leer_excel(ruta, HOJAS_PANEL, COLUMNAS_PANEL, ('nombre', 'estado', 'fecha'))


def leer_tracker(ruta):
    return leer_excel(ruta, HOJAS_TRACKER, COLUMNAS_TRACKER, ('nombre', 'rating'))


def agrupar(filas, asociador, origen):
    """({id: [filas]}, [nombres sin asociar])."""
    por_id, sin_asociar = {}, []
    for f in filas:
        oid, motivo = asociador.buscar(f['nombre'])
        if oid:
            por_id.setdefault(oid, []).append(f)
        else:
            sin_asociar.append(f'{origen}, hoja «{f["hoja"]}», fila {f["fila"]}: {motivo}')
    return por_id, sin_asociar


def rango_valoracion(valor):
    """Número para ordenar por valoración (mayor = mejor); None si no se reconoce."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip()
    if not texto:
        return None
    if set(texto) <= set('★☆* '):
        return float(texto.count('★') + texto.count('*'))
    try:
        return float(texto.replace(',', '.'))
    except ValueError:
        return PALABRAS_VALORACION.get(normalizar(texto))


def texto_valoracion(valor):
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return '' if valor is None else str(valor).strip()


def cruzar(exp, orgs, filas_panel=(), filas_tracker=()):
    """Reúne para cada expediente los datos de organizaciones.json, el panel y el tracker.

    Devuelve (fichas, huerfanos, avisos): fichas es la lista de diccionarios con
    los datos de cada ficha, en el orden del expediente; huerfanos, los id del
    expediente sin organización; avisos, los nombres sin asociar y las empresas
    con varias filas en el panel.
    """
    asociador = Asociador(orgs, exp['empresas'])
    panel, sin_panel = agrupar(filas_panel, asociador, 'Panel')
    tracker, sin_tracker = agrupar(filas_tracker, asociador, 'Tracker')
    por_id = {o['id']: o for o in orgs}
    fichas, huerfanos, avisos = [], [], sin_panel + sin_tracker
    for oid in list(panel) + list(tracker):
        if oid not in exp['empresas']:
            avisos.append(f'«{oid}» aparece en el panel o en el tracker pero no tiene expediente.')
    for oid, e in exp['empresas'].items():
        org = por_id.get(oid)
        if org is None:
            huerfanos.append(oid)
            continue
        filas = panel.get(oid, [])
        if len(filas) > 1:
            avisos.append(f'Panel: {org["nombre"]} aparece en {len(filas)} filas '
                          f'({", ".join(f"{f["hoja"]} {f["fila"]}" for f in filas)}); se usa la primera.')
        fila = filas[0] if filas else {}
        valoraciones = [f['rating'] for f in tracker.get(oid, []) if not vacio(f['rating'])]
        mejor = max(valoraciones, key=lambda v: (rango_valoracion(v) is not None, rango_valoracion(v) or 0),
                    default=None)
        fichas.append({
            'id': oid, 'org': org, 'exp': e,
            'estado': '' if fila.get('estado') is None else str(fila['estado']).strip(),
            'proximo': fecha_es(fila.get('fecha')),
            'valoracion': texto_valoracion(mejor),
            'rango': rango_valoracion(mejor),
        })
    return fichas, huerfanos, list(dict.fromkeys(avisos))


# --- Word: piezas básicas ------------------------------------------------------

PAG_ALTO, PAG_ANCHO, MARGEN_SUP, MARGEN_INF, MARGEN_LAT = 16838, 11906, 850, 850, 1020
ANCHO = PAG_ANCHO - 2 * MARGEN_LAT  # 9866
ALTO_RENGLON = 454
ALTO_ETIQUETA = 320
HUECO = 240
RENGLONES_NOTAS = (  # recuadros de la página de notas: títulos y renglones
    (('FECHA Y HORA', 'LUGAR O MEDIO'), 2),
    (('ASISTENTES DEL CLIENTE', 'ASISTENTES POR NUESTRA PARTE'), 3),
    (('NOTAS',), 12),
    (('ACUERDOS',), 4),
    (('PRÓXIMOS PASOS',), 4),
)

CAL = '<w:rFonts w:ascii="Calibri" w:cs="Calibri" w:eastAsia="Calibri" w:hAnsi="Calibri"/>'
NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
TIPO_ENLACE = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink'


def rpr(sz, b=False, i=False, color='000000'):
    s = CAL
    if b:
        s += '<w:b/><w:bCs/>'
    if i:
        s += '<w:i/><w:iCs/>'
    return f'<w:rPr>{s}<w:color w:val="{color}"/><w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/></w:rPr>'


R_TEXTO = rpr(18)
R_NEGRITA = rpr(18, b=True)
R_ETIQUETA = rpr(14, b=True)
R_GRIS = rpr(15, color='595959')
TAB = '<w:r><w:tab/></w:r>'


def run(texto, props=R_TEXTO):
    return f'<w:r>{props}<w:t xml:space="preserve">{escape(CONTROL.sub("", str(texto)))}</w:t></w:r>'


def ppr(mantener=False, salto=False, borde=None, tabs=(), antes=None, despues=None, linea=None, sangria=None,
        jc=None):
    """Propiedades de párrafo, en el orden que exige Word."""
    s = '<w:keepNext/>' if mantener else ''
    if salto:
        s += '<w:pageBreakBefore/>'
    if borde:
        s += f'<w:pBdr><w:bottom w:val="single" w:sz="{borde[1]}" w:space="1" w:color="{borde[0]}"/></w:pBdr>'
    if tabs:
        s += '<w:tabs>' + ''.join(f'<w:tab w:val="{v}" w:pos="{p}"/>' for v, p in tabs) + '</w:tabs>'
    if antes is not None or despues is not None or linea:
        atr = f' w:before="{antes}"' if antes is not None else ''
        atr += f' w:after="{despues}"' if despues is not None else ''
        atr += f' w:line="{linea}" w:lineRule="exact"' if linea else ''
        s += f'<w:spacing{atr}/>'
    if sangria:
        s += f'<w:ind w:left="{sangria}" w:hanging="{sangria}"/>'
    if jc:
        s += f'<w:jc w:val="{jc}"/>'
    return s


def parrafo(runs='', propiedades=''):
    return f'<w:p><w:pPr>{propiedades}</w:pPr>{runs}</w:p>' if propiedades else f'<w:p>{runs}</w:p>'


def p(texto='', props=R_TEXTO, **kw):
    return parrafo(run(texto, props) if texto else '', ppr(**kw))


def parrafos(texto, props=R_TEXTO):
    """Un párrafo por línea del texto."""
    return ''.join(p(linea.strip(), props) for linea in str(texto).splitlines() if linea.strip())


def vineta(texto, props=R_TEXTO):
    return parrafo(run('•', props) + TAB + run(texto, props), ppr(tabs=(('left', 180),), sangria=180))


def separador(alto):
    return parrafo('', ppr(antes=0, despues=0, linea=alto))


SALTO = parrafo('', ppr(salto=True, antes=0, despues=0, linea=20))
FINAL = separador(20)

GRIS_LINEA = ('d8d8d8', 4)
GRIS_RENGLON = ('cfcfcf', 4)
NEGRO = ('000000', 8)
NEGRO_FINO = ('000000', 4)
SIN = (None, None, None, None)
INFERIOR = (None, None, GRIS_LINEA, None)


def borde(lado, spec):
    if spec is None:
        return f'<w:{lado} w:val="nil"/>'
    color, sz = spec
    return f'<w:{lado} w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'


def celda(contenido, bordes=SIN, mar=(50, 100, 50, 100), relleno=None, valign=None, span=1):
    return {'xml': contenido, 'bordes': bordes, 'mar': mar, 'relleno': relleno, 'valign': valign, 'span': span}


def fila(celdas, alto=None, exacto=False, cabecera=False, partir=True):
    return {'celdas': celdas, 'alto': alto, 'exacto': exacto, 'cabecera': cabecera, 'partir': partir}


def tabla(columnas, filas):
    """Tabla de primer nivel. Las fichas nunca anidan tablas: desbordan en Word."""
    grid = ''.join(f'<w:gridCol w:w="{c}"/>' for c in columnas)
    sin_bordes = ''.join(f'<w:{lado} w:val="nil"/>'
                         for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'))
    xml = (f'<w:tbl><w:tblPr><w:tblStyle w:val="TablaFicha"/><w:tblW w:w="{sum(columnas)}" w:type="dxa"/>'
           f'<w:jc w:val="left"/><w:tblBorders>{sin_bordes}</w:tblBorders><w:tblLayout w:type="fixed"/>'
           f'<w:tblLook w:val="0000"/></w:tblPr><w:tblGrid>{grid}</w:tblGrid>')
    for f in filas:
        trpr = '' if f['partir'] else '<w:cantSplit/>'
        if f['alto']:
            trpr += f'<w:trHeight w:val="{f["alto"]}" w:hRule="{"exact" if f["exacto"] else "atLeast"}"/>'
        if f['cabecera']:
            trpr += '<w:tblHeader/>'
        xml += '<w:tr>' + (f'<w:trPr>{trpr}</w:trPr>' if trpr else '')
        i = 0
        for c in f['celdas']:
            ancho = sum(columnas[i:i + c['span']])
            i += c['span']
            tcpr = f'<w:tcW w:w="{ancho}" w:type="dxa"/>'
            if c['span'] > 1:
                tcpr += f'<w:gridSpan w:val="{c["span"]}"/>'
            tcpr += '<w:tcBorders>' + ''.join(borde(lado, b) for lado, b in
                                               zip(('top', 'left', 'bottom', 'right'), c['bordes'])) + '</w:tcBorders>'
            if c['relleno']:
                tcpr += f'<w:shd w:val="clear" w:color="auto" w:fill="{c["relleno"]}"/>'
            tcpr += '<w:tcMar>' + ''.join(f'<w:{lado} w:w="{v}" w:type="dxa"/>' for lado, v in
                                          zip(('top', 'left', 'bottom', 'right'), c['mar'])) + '</w:tcMar>'
            if c['valign']:
                tcpr += f'<w:vAlign w:val="{c["valign"]}"/>'
            xml += f'<w:tc><w:tcPr>{tcpr}</w:tcPr>{c["xml"] or parrafo()}</w:tc>'
        if i != len(columnas):
            raise ValueError(f'fila de {i} columnas en una tabla de {len(columnas)}')
        xml += '</w:tr>'
    return xml + '</w:tbl>'


def repartir(anchos):
    """Escala los anchos de columna al ancho útil de la página."""
    total = sum(anchos)
    escalados = [round(a * ANCHO / total) for a in anchos]
    escalados[-1] += ANCHO - sum(escalados)
    return escalados


def tabla_datos(columnas, titulos, filas, centradas=()):
    """Tabla con fila de títulos (se repite en cada página) y una fila por elemento.

    Cada fila es una lista con el XML (párrafos) de cada celda.
    """
    cab = fila([celda(p(t, R_ETIQUETA, jc='center' if i in centradas else None), (None, None, NEGRO_FINO, None),
                      mar=(40, 80, 40, 80)) for i, t in enumerate(titulos)], cabecera=True, partir=False)
    cuerpo = [fila([celda(x, INFERIOR, mar=(50, 80, 50, 80)) for x in f]) for f in filas]
    return tabla(columnas, [cab, *cuerpo])


def titulo_seccion(texto, derecha=''):
    runs = run(texto, rpr(16, b=True))
    if derecha:
        runs += TAB + run(derecha, R_GRIS)
    return parrafo(runs, ppr(mantener=True, borde=NEGRO, tabs=(('right', ANCHO),), antes=320, despues=100))


def subtitulo(texto):
    return p(texto, R_ETIQUETA, mantener=True, antes=160, despues=60)


def escribir_docx(ruta, cuerpo, pie, enlaces=()):
    """Escribe el .docx con las partes fijas de la plantilla (estilos y pie del bloc)."""
    seccion = (f'<w:sectPr><w:footerReference w:type="default" r:id="rId6"/>'
               f'<w:pgSz w:w="{PAG_ANCHO}" w:h="{PAG_ALTO}"/>'
               f'<w:pgMar w:top="{MARGEN_SUP}" w:right="{MARGEN_LAT}" w:bottom="{MARGEN_INF}" '
               f'w:left="{MARGEN_LAT}" w:header="708" w:footer="400" w:gutter="0"/>'
               f'<w:pgNumType w:start="1"/></w:sectPr>')
    xml = (f"<?xml version='1.0' encoding='UTF-8' standalone='yes'?>\n"
           f'<w:document {NS}><w:body>{cuerpo}{FINAL}{seccion}</w:body></w:document>')
    rels = (PLANTILLA / 'word' / '_rels' / 'document.xml.rels').read_text(encoding='utf-8')
    rels = rels.replace('</Relationships>', ''.join(
        f'<Relationship Id="{rid}" Type="{TIPO_ENLACE}" Target={quoteattr(destino)} TargetMode="External"/>'
        for rid, destino in enlaces) + '</Relationships>')
    pie_xml = (PLANTILLA / 'word' / 'footer1.xml').read_text(encoding='utf-8').replace('{PIE}', escape(pie))
    with zipfile.ZipFile(ruta, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', (PLANTILLA / '[Content_Types].xml').read_bytes())
        z.writestr('_rels/.rels', (PLANTILLA / '_rels' / '.rels').read_bytes())
        z.writestr('word/document.xml', xml.encode('utf-8'))
        z.writestr('word/_rels/document.xml.rels', rels.encode('utf-8'))
        for parte in ('word/settings.xml', 'word/fontTable.xml', 'word/numbering.xml', 'word/styles.xml',
                      'word/theme/theme1.xml'):
            z.writestr(parte, (PLANTILLA / parte).read_bytes())
        z.writestr('word/footer1.xml', pie_xml.encode('utf-8'))
    return ruta


# --- Ficha maestra -------------------------------------------------------------

def etiqueta_grupo(org):
    """«Grupo A · Promotora»"""
    grupo = f"Grupo {org.get('grupo', '')}"
    return f"{grupo} · {org['subtipo']}" if org.get('subtipo') else grupo


def orden_prioridad(contactos):
    """Interlocutores primero, después asistencia confirmada y el resto (como el bloc)."""
    return sorted(contactos, key=lambda c: (not c.get('interlocutor'), not c.get('confirmado_feria')))


def recientes_primero(entradas):
    """De la fecha más reciente a la más antigua; con la misma fecha, en el orden en que se añadieron."""
    return sorted(sorted(entradas, key=lambda x: x['id']), key=lambda x: x['fecha'], reverse=True)


def fuente_y_fecha(x):
    return ' · '.join(t for t in (x.get('fuente', ''), fecha_es(x.get('fecha'))) if t)


def caja_letra(org, mar=(100, 60, 100, 60)):
    return celda(p(org.get('grupo', ''), rpr(36, b=True, color='ffffff'), jc='center'), mar=mar,
                 relleno='000000', valign='center')


def cabecera_ficha(d, fecha_gen):
    org = d['org']
    titulo = (p('FICHA MAESTRA', rpr(13, b=True, color='595959'))
              + p(org['nombre'], rpr(32, b=True)) + p(etiqueta_grupo(org)))
    celdas = [caja_letra(org), celda(titulo, mar=(100, 160, 100, 100), relleno='f1f1f1', valign='center')]
    columnas = [900, ANCHO - 900]
    if d['valoracion']:
        derecha = p('VALORACIÓN', rpr(13, b=True), jc='right') + p(d['valoracion'], rpr(30, b=True), jc='right')
        celdas.append(celda(derecha, mar=(80, 100, 80, 160), relleno='f1f1f1', valign='center'))
        columnas = [900, ANCHO - 900 - 2400, 2400]
    datos = [(etq, valor, ancho) for etq, valor, ancho in
             (('ESTADO', d['estado'], 4000), ('PRÓXIMO COMPROMISO', d['proximo'], 2400),
              ('FICHA GENERADA', fecha_es(fecha_gen), 2000)) if valor]
    tira = [celda(p(etq, R_ETIQUETA) + p(valor, R_NEGRITA), (None, None, NEGRO, None), mar=(60, 100, 80, 100))
            for etq, valor, _ in datos]
    return (tabla(columnas, [fila(celdas, partir=False)]) + separador(120)
            + tabla(repartir([a for _, _, a in datos]), [fila(tira, partir=False)]))


def seccion_situacion(e):
    s = e['situacion_actual']
    if not s:
        return ''
    caja = celda(parrafos(s['texto']), (NEGRO, None, NEGRO, None), mar=(100, 140, 100, 140), relleno='f1f1f1')
    return (titulo_seccion('SITUACIÓN ACTUAL', f'Actualizada el {fecha_es(s["fecha"])}')
            + tabla([ANCHO], [fila([caja])]))


def seccion_contactos(org, fecha_org):
    contactos = orden_prioridad(org.get('contactos') or [])
    if not contactos and not org.get('telefono_empresa'):
        return ''
    filas = []
    for c in contactos:
        persona = p(c.get('nombre', ''), R_NEGRITA if c.get('interlocutor') else R_TEXTO)
        notas = (['Interlocutor'] if c.get('interlocutor') else []) + ([c['nota']] if c.get('nota') else [])
        if notas:
            persona += p(' · '.join(notas), R_GRIS)
        filas.append([persona, p(c.get('puesto', '')), p(c.get('email', '')), p(c.get('telefono', ''))])
    if org.get('telefono_empresa'):
        etiqueta = 'Tel. organización' if org.get('subtipo') in SUBTIPOS_ORGANIZACION else 'Tel. empresa'
        filas.append([p(etiqueta, R_GRIS), '', '', p(org['telefono_empresa'])])
    return (titulo_seccion('CONTACTOS', f'Copia de organizaciones.json del {fecha_es(fecha_org)}')
            + tabla_datos([2400, 2466, 3100, 1900], ['PERSONA', 'CARGO', 'CORREO', 'TELÉFONO'], filas))


def parrafo_proyecto(pr):
    """«Nombre — Ubicación · Año · Estado», como en el bloc."""
    cola = f" — {pr['ubicacion']}" if pr.get('ubicacion') else ''
    cola += f" · {pr['anio']}" if pr.get('anio') else ''
    cola += ' · ' if pr.get('estado') else ''
    runs = run('•', R_TEXTO) + TAB + run(pr.get('nombre', ''), R_NEGRITA)
    if cola:
        runs += run(cola)
    if pr.get('estado'):
        runs += run(pr['estado'], rpr(18, i=True))
    return parrafo(runs, ppr(tabs=(('left', 180),), sangria=180))


def seccion_perfil(org, e):
    filas = []
    if org.get('perfil'):
        filas.append(('PERFIL', parrafos(org['perfil'])))
    proyectos = org.get('proyectos') or []
    conocidos = org.get('proyectos_conocidos')
    xml = ''.join(parrafo_proyecto(pr) for pr in proyectos)
    if isinstance(conocidos, int) and conocidos > len(proyectos):
        xml += p(f'Proyectos conocidos: {conocidos}', R_GRIS)
    if xml:
        filas.append(('PROYECTOS', xml))
    trabaja = org.get('trabaja_con')
    if trabaja:
        filas.append(('CON QUIÉN TRABAJA', p(', '.join(trabaja) if isinstance(trabaja, list) else trabaja)))
    if org.get('uso_gfs') in ('Sí', 'No'):
        filas.append(('USO DE GFS', p(org['uso_gfs'], R_NEGRITA)))
    xml = tabla([1900, ANCHO - 1900], [fila([celda(p(etq, R_ETIQUETA), INFERIOR, mar=(70, 0, 70, 100)),
                                             celda(contenido, INFERIOR, mar=(70, 100, 70, 0))])
                                       for etq, contenido in filas]) if filas else ''
    if e['proyectos']:
        filas_exp = [[p(x['nombre'], R_NEGRITA), p(x.get('ubicacion', '')),
                      p(x.get('situacion', '')) + p(x['fuente'], R_GRIS), p(fecha_es(x['fecha']))]
                     for x in recientes_primero(e['proyectos'])]
        xml += subtitulo('PROYECTOS EN SEGUIMIENTO') + tabla_datos(
            [2500, 1800, 4266, 1300], ['PROYECTO', 'UBICACIÓN', 'SITUACIÓN', 'FECHA'], filas_exp)
    return titulo_seccion('PERFIL Y PROYECTOS') + xml if xml else ''


def seccion_lo_que_sabemos(e):
    if not e['lo_que_sabemos']:
        return ''
    orden = list(NATURALEZAS)
    entradas = sorted(recientes_primero(e['lo_que_sabemos']), key=lambda x: orden.index(x['naturaleza']))
    filas = [fila([celda(p(NATURALEZAS[x['naturaleza']], R_ETIQUETA), INFERIOR, mar=(60, 0, 60, 100)),
                   celda(parrafos(x['texto']) + p(fuente_y_fecha(x), R_GRIS), INFERIOR, mar=(60, 100, 60, 0))])
             for x in entradas]
    return titulo_seccion('LO QUE SABEMOS') + tabla([1500, ANCHO - 1500], filas)


COLUMNAS_TANQUES = (  # campo, título, ancho relativo, centrada
    ('proyecto', 'PROYECTO', 1500, False), ('uso', 'USO', 1200, False), ('cantidad', 'CANT.', 600, True),
    ('diametro_m', 'Ø (m)', 750, True), ('altura_m', 'ALTURA (m)', 850, True),
    ('volumen_m3', 'VOLUMEN (m³)', 1000, True), ('techo', 'TECHO', 1100, False),
    ('aislamiento', 'AISLAMIENTO', 1100, False), ('otros', 'OTROS', 1400, False),
)


def seccion_tanques(e):
    if not e['tanques']:
        return ''
    tanques = recientes_primero(e['tanques'])
    columnas = [c for c in COLUMNAS_TANQUES if any(not vacio(t.get(c[0])) for t in tanques)]
    centradas = [i for i, c in enumerate(columnas) if c[3]]
    filas = []
    for t in tanques:
        valores = [numero_es(t.get(c[0])) if c[0] in ('cantidad', 'diametro_m', 'altura_m', 'volumen_m3')
                   else t.get(c[0], '') for c in columnas]
        filas.append([p(v, jc='center' if i in centradas else None) for i, v in enumerate(valores)]
                     + [p(t['fuente'], R_GRIS) + p(fecha_es(t['fecha']), R_GRIS)])
    return titulo_seccion('TANQUES') + tabla_datos(
        repartir([c[2] for c in columnas] + [1400]), [c[1] for c in columnas] + ['FUENTE'], filas, centradas)


def seccion_cotizaciones(e):
    if not e['cotizaciones']:
        return ''
    cotizaciones = recientes_primero(e['cotizaciones'])
    con_importe = any(x.get('importe_eur') is not None for x in cotizaciones)
    filas = []
    for x in cotizaciones:
        f = [p(fecha_es(x['fecha'])), p(x['descripcion']), p(x['estado'])]
        if con_importe:
            importe = x.get('importe_eur')
            f.append(p(f'{numero_es(importe)} €' if importe is not None else '', jc='right'))
        filas.append(f + [p(x['fuente'], R_GRIS)])
    titulos = ['FECHA', 'DESCRIPCIÓN', 'ESTADO'] + (['IMPORTE'] if con_importe else []) + ['FUENTE']
    anchos = [1150, 3900, 1500] + ([1400] if con_importe else []) + [1900]
    return titulo_seccion('COTIZACIONES') + tabla_datos(repartir(anchos), titulos, filas)


def seccion_pendientes(e):
    abiertos = [x for x in e['pendientes'] if x['estado'] == 'abierto']
    if not abiertos:
        return ''
    abiertos.sort(key=lambda x: (x.get('fecha_limite') is None, x.get('fecha_limite') or '', x['id']))
    filas = [[p(x['accion']), p(DE_QUIEN[x['de_quien']]), p(fecha_es(x.get('fecha_limite')))] for x in abiertos]
    return titulo_seccion('PENDIENTES ABIERTOS') + tabla_datos(
        [6066, 1800, 2000], ['ACCIÓN', 'DE QUIÉN', 'FECHA LÍMITE'], filas)


def seccion_competencia(e):
    if not e['competencia']:
        return ''
    filas = [fila([celda(p(fecha_es(x['fecha'])), INFERIOR, mar=(60, 0, 60, 100)),
                   celda(parrafos(x['texto']) + p(x['fuente'], R_GRIS), INFERIOR, mar=(60, 100, 60, 0))])
             for x in recientes_primero(e['competencia'])]
    return titulo_seccion('COMPETENCIA Y OBJECIONES') + tabla([1500, ANCHO - 1500], filas)


def orden_cronologia(entradas):
    """De la más reciente a la más antigua."""
    return sorted(entradas, key=lambda x: (x['fecha'], x.get('hora') or '', x['id']), reverse=True)


def seccion_cronologia(e):
    if not e['cronologia']:
        return ''
    filas = []
    for x in orden_cronologia(e['cronologia']):
        izquierda = p(fecha_es(x['fecha']), R_NEGRITA)
        if x.get('hora'):
            izquierda += p(x['hora'])
        izquierda += p(TIPOS[x['tipo']])
        if x['origen'] == 'automatico':
            izquierda += p('AUTOMÁTICA', R_ETIQUETA, antes=40)
        derecha = ''
        if x.get('participantes'):
            derecha += p('Participantes: ' + ', '.join(x['participantes']), R_GRIS)
        derecha += parrafos(x['resumen'])
        for titulo, campo in (('Datos clave', 'datos_clave'), ('Acuerdos', 'acuerdos')):
            if x.get(campo):
                derecha += p(titulo.upper(), R_ETIQUETA, antes=60) + ''.join(vineta(t) for t in x[campo])
        filas.append(fila([celda(izquierda, INFERIOR, mar=(70, 0, 70, 100)),
                           celda(derecha, INFERIOR, mar=(70, 100, 70, 0))]))
    return titulo_seccion('CRONOLOGÍA') + tabla([1700, ANCHO - 1700], filas)


def seccion_documentos(e):
    xml = ''
    if e['documentos']:
        filas = [[p(fecha_es(x['fecha'])), p(x['nombre']), p(x.get('ruta', ''), R_GRIS)]
                 for x in recientes_primero(e['documentos'])]
        xml += subtitulo('DOCUMENTOS') + tabla_datos([1150, 3800, 4916], ['FECHA', 'DOCUMENTO', 'RUTA'], filas)
    con_ref = [x for x in orden_cronologia(e['cronologia']) if x['fuente'].get('ref', '').strip()]
    if con_ref:
        filas = [[p(fecha_es(x['fecha'])), p(TIPOS[x['tipo']]), p(FUENTES[x['fuente']['tipo']]),
                  p(x['fuente']['ref'], R_GRIS)] for x in con_ref]
        xml += subtitulo('REFERENCIAS DE LA CRONOLOGÍA') + tabla_datos(
            [1150, 1900, 1600, 5216], ['FECHA', 'ENTRADA', 'FUENTE', 'REFERENCIA'], filas)
    return titulo_seccion('DOCUMENTOS Y FUENTES') + xml if xml else ''


def recuadros(titulos, renglones):
    """Uno o dos recuadros (lado a lado) con su título y renglones para escribir."""
    if len(titulos) == 1:
        columnas = [ANCHO]
    else:
        mitad = (ANCHO - HUECO) // 2
        columnas = [mitad, HUECO, ANCHO - HUECO - mitad]
    filas = []
    for k in range(renglones + 1):
        celdas = []
        for j, titulo in enumerate(titulos):
            if j:
                celdas.append(celda('', SIN, mar=(0, 0, 0, 0)))
            abajo = NEGRO if k == renglones else (GRIS_RENGLON if k else None)
            bordes = (NEGRO if k == 0 else None, NEGRO, abajo, NEGRO)
            contenido = p(titulo, R_ETIQUETA) if k == 0 else ''
            celdas.append(celda(contenido, bordes, mar=(70, 120, 0, 120) if k == 0 else (0, 120, 0, 120)))
        filas.append(fila(celdas, alto=ALTO_ETIQUETA if k == 0 else ALTO_RENGLON, exacto=True, partir=False))
    return tabla(columnas, filas)


def pagina_notas(org):
    cab = tabla([900, ANCHO - 900], [fila([
        caja_letra(org),
        celda(p('NOTAS DE REUNIÓN', rpr(13, b=True, color='595959')) + p(org['nombre'], rpr(28, b=True)),
              mar=(100, 160, 100, 100), relleno='f1f1f1', valign='center'),
    ], partir=False)])
    partes = [SALTO, cab]
    for titulos, renglones in RENGLONES_NOTAS:
        partes += [separador(200), recuadros(titulos, renglones)]
    return ''.join(partes)


def cuerpo_ficha(d, fecha_gen, fecha_org):
    org, e = d['org'], d['exp']
    return ''.join((
        cabecera_ficha(d, fecha_gen),
        seccion_situacion(e),
        seccion_contactos(org, fecha_org),
        seccion_perfil(org, e),
        seccion_lo_que_sabemos(e),
        seccion_tanques(e),
        seccion_cotizaciones(e),
        seccion_pendientes(e),
        seccion_competencia(e),
        seccion_cronologia(e),
        seccion_documentos(e),
        pagina_notas(org),
    ))


def escribir_ficha(d, ruta, fecha_gen, fecha_org):
    return escribir_docx(ruta, cuerpo_ficha(d, fecha_gen, fecha_org), f'Ficha maestra · {d["org"]["nombre"]}')


# --- Índice --------------------------------------------------------------------

def ultima_interaccion(e):
    return orden_cronologia(e['cronologia'])[0] if e['cronologia'] else None


def orden_indice(fichas):
    """Por valoración (de mayor a menor) y, después, por la última interacción (más reciente primero)."""
    def clave(d):
        ultima = ultima_interaccion(d['exp'])
        dia = datetime.date.fromisoformat(ultima['fecha']).toordinal() if ultima else 0
        return (d['rango'] is None, -(d['rango'] or 0), ultima is None, -dia, normalizar(d['org']['nombre']))
    return sorted(fichas, key=clave)


def nombre_ficha(oid, extension):
    return f'Ficha_{oid}.{extension}'


def escribir_indice(fichas, ruta, fecha_gen):
    portada = tabla([ANCHO], [fila([celda(
        p('STORA', rpr(22, b=True, color='c5c5c5'), despues=120)
        + p('Fichas maestras', rpr(48, b=True, color='ffffff'), despues=80)
        + p(f'Índice · {fecha_es(fecha_gen)} · {len(fichas)} empresa{"s" if len(fichas) != 1 else ""}',
            rpr(22, color='ffffff')),
        mar=(360, 320, 380, 320), relleno='000000')], partir=False)])
    filas, enlaces = [], []
    for n, d in enumerate(orden_indice(fichas), 1):
        rid = f'rIdFicha{n}'
        enlaces.append((rid, f'{CARPETA_FICHAS}/{nombre_ficha(d["id"], "pdf")}'))
        nombre = f'<w:hyperlink r:id="{rid}" w:history="1">{run(d["org"]["nombre"], R_NEGRITA)}</w:hyperlink>'
        ultima = ultima_interaccion(d['exp'])
        abiertos = sum(1 for x in d['exp']['pendientes'] if x['estado'] == 'abierto')
        org = d['org']
        tipo = ' · '.join(t for t in (org.get('grupo'), org.get('subtipo')) if t)
        filas.append([
            parrafo(nombre), p(tipo), p(d['valoracion'], jc='center'),
            p(d['estado']) + (p(f'Próximo: {d["proximo"]}', R_GRIS) if d['proximo'] else ''),
            (p(fecha_es(ultima['fecha'])) + p(TIPOS[ultima['tipo']], R_GRIS)) if ultima else '',
            p(str(abiertos), jc='center'),
        ])
    cuerpo = portada + separador(240) + tabla_datos(
        [2700, 1500, 1050, 2216, 1500, 900],
        ['EMPRESA', 'TIPO', 'VALORACIÓN', 'ESTADO', 'ÚLTIMA INTERACCIÓN', 'PENDIENTES'], filas, centradas=(2, 5))
    return escribir_docx(ruta, cuerpo, 'Fichas maestras · Índice', enlaces)


# --- PDF -----------------------------------------------------------------------

RUTAS_LIBREOFFICE = (
    '/Applications/LibreOffice.app/Contents/MacOS/soffice',
    r'C:\Program Files\LibreOffice\program\soffice.exe',
    r'C:\Program Files (x86)\LibreOffice\program\soffice.exe',
)
FILTRO_ENLACES_RELATIVOS = 'pdf:writer_pdf_Export:{"ExportLinksRelativeFsys":{"type":"boolean","value":"true"}}'


def soffice():
    for nombre in ('soffice', 'libreoffice'):
        ruta = shutil.which(nombre)
        if ruta:
            return ruta
    for ruta in RUTAS_LIBREOFFICE:
        if Path(ruta).exists():
            return ruta
    raise ErrorGeneracion('No se encuentra LibreOffice (soffice). Es necesario para generar el PDF.')


def convertir_a_pdf(docxs, carpeta, enlaces_relativos=False):
    """Convierte los .docx a PDF con una sola ejecución de LibreOffice (mismas
    opciones que scripts/bloc_pdf.py) y devuelve las rutas de los PDF."""
    docxs, carpeta = [Path(d) for d in docxs], Path(carpeta)
    if not docxs:
        return []
    perfil = Path(tempfile.mkdtemp(prefix='stora_lo_'))
    try:
        orden = [soffice(), f'-env:UserInstallation={perfil.as_uri()}', '--headless', '--convert-to',
                 FILTRO_ENLACES_RELATIVOS if enlaces_relativos else 'pdf', '--outdir', str(carpeta),
                 *map(str, docxs)]
        r = subprocess.run(orden, capture_output=True, text=True, timeout=120 + 30 * len(docxs))
    finally:
        shutil.rmtree(perfil, ignore_errors=True)
    pdfs = [carpeta / (d.stem + '.pdf') for d in docxs]
    faltan = [pdf.name for pdf in pdfs if not pdf.exists()]
    if r.returncode != 0 or faltan:
        raise ErrorGeneracion(f'LibreOffice no pudo convertir {", ".join(faltan) or "los documentos"}:\n'
                              f'{r.stdout}\n{r.stderr}')
    return pdfs


def hay_pdfplumber():
    return importlib.util.find_spec('pdfplumber') is not None


LIMITE_CABECERA = 130  # pt desde el borde superior: zona en la que se busca la cabecera de cada página


def analizar_pdf(ruta_pdf):
    """Localiza en el PDF el comienzo de cada ficha y cada página de notas.

    Mismo método que la comprobación de scripts/validar.py --pdf: una cabecera
    se reconoce por la letra del grupo, en cuerpo grande, dentro del recuadro
    negro de la esquina superior izquierda, y su rótulo dice si es el comienzo
    de la ficha («FICHA MAESTRA») o la página de notas («NOTAS DE REUNIÓN»).
    Ningún texto de los datos se imprime en ese cuerpo, así que no puede
    confundirse con una cabecera. Devuelve (total_paginas, inicios, notas,
    fuera): inicios y notas son las páginas (índice 0) con cada cabecera; fuera,
    los textos que rebasan el margen derecho.
    """
    import pdfplumber

    inicios, notas, fuera = [], [], []
    limite_derecho = (PAG_ANCHO - MARGEN_LAT) / 20 + 3
    with pdfplumber.open(str(ruta_pdf)) as doc:
        total = len(doc.pages)
        for n, pagina in enumerate(doc.pages):
            palabras = pagina.extract_words(extra_attrs=['size'])
            arriba = [w for w in palabras if w['top'] < LIMITE_CABECERA]
            if any(w['size'] > 15 and w['x0'] < 100 for w in arriba):
                rotulo = normalizar(' '.join(w['text'] for w in arriba))
                if 'notas de reunion' in rotulo:
                    notas.append(n)
                elif 'ficha maestra' in rotulo:
                    inicios.append(n)
            fuera += [f'página {n + 1}: «{w["text"]}» rebasa el margen derecho'
                      for w in palabras if w['x1'] > limite_derecho]
    return total, inicios, notas, fuera


def problemas_paginacion(ruta_pdf):
    """Cada ficha empieza en página nueva y termina con su página de notas, que
    ocupa una sola página; ningún texto se sale de la página."""
    total, inicios, notas, fuera = analizar_pdf(ruta_pdf)
    problemas = []
    if not inicios or inicios[0] != 0:
        problemas.append('la ficha no empieza en la primera página')
    for k, inicio in enumerate(inicios):
        fin = inicios[k + 1] if k + 1 < len(inicios) else total
        propias = [n for n in notas if inicio <= n < fin]
        if len(propias) != 1:
            problemas.append(f'la ficha {k + 1} tiene {len(propias)} páginas de notas (debe tener una)')
        elif propias[0] != fin - 1:
            problemas.append(f'la página de notas de la ficha {k + 1} ocupa {fin - propias[0]} páginas')
        elif propias[0] == inicio:
            problemas.append(f'la ficha {k + 1} no tiene contenido antes de la página de notas')
    return total, problemas + fuera


# --- generar -------------------------------------------------------------------

def huella_formato():
    """Huella del código y la plantilla: si cambian, se regeneran todas las fichas."""
    h = hashlib.sha256(Path(__file__).read_bytes())
    for parte in sorted(PLANTILLA.rglob('*')):
        if parte.is_file():
            h.update(parte.relative_to(PLANTILLA).as_posix().encode())
            h.update(parte.read_bytes())
    return h.hexdigest()


def huella(d, formato):
    contenido = {'formato': formato, 'organizacion': d['org'], 'expediente': d['exp'], 'estado': d['estado'],
                 'proximo': d['proximo'], 'valoracion': d['valoracion']}
    return hashlib.sha256(json.dumps(contenido, sort_keys=True, ensure_ascii=False, default=str)
                          .encode('utf-8')).hexdigest()


def cargar_estado(ruta):
    try:
        estado = leer_json(ruta)
    except (OSError, json.JSONDecodeError):
        return {'version': VERSION, 'fichas': {}}
    return estado if isinstance(estado.get('fichas'), dict) else {'version': VERSION, 'fichas': {}}


def generar(ruta_exp, salida, ruta_org, ruta_panel, ruta_tracker, todas=False, fecha=None, fecha_org=None,
            informe=print):
    fecha = fecha or datetime.date.today().isoformat()
    salida = Path(salida)
    exp, errores = cargar_expedientes(ruta_exp)
    if errores:
        for m in errores:
            informe(f'ERROR      {m}')
        informe('\nExpediente no válido: no se genera ninguna ficha.')
        return 1
    try:
        orgs = cargar_organizaciones(ruta_org)
        filas_panel, avisos_panel = leer_panel(ruta_panel)
        filas_tracker, avisos_tracker = leer_tracker(ruta_tracker)
    except ErrorDatos as exc:
        informe(f'ERROR      {exc}')
        return 1
    fecha_org = fecha_organizaciones(ruta_org, fecha_org)
    fichas, huerfanos, avisos = cruzar(exp, orgs, filas_panel, filas_tracker)
    avisos = avisos_panel + avisos_tracker + avisos
    for oid in huerfanos:
        informe(f'AVISO      expediente «{oid}» sin correspondencia en {Path(ruta_org).name}: no se genera su ficha.')
    if avisos:
        informe(f'AVISO      {len(avisos)} aviso(s) del panel o del tracker (detalle con «validar»).')

    carpeta = salida / CARPETA_FICHAS
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta_estado = salida / ESTADO
    estado = cargar_estado(ruta_estado)
    formato = huella_formato()
    pendientes = []
    for d in fichas:
        h = huella(d, formato)
        anterior = estado['fichas'].get(d['id'], {})
        if todas or anterior.get('huella') != h or not all(
                (carpeta / nombre_ficha(d['id'], x)).exists() for x in ('docx', 'pdf')):
            pendientes.append((d, h))

    comprobar = hay_pdfplumber()
    if not comprobar:
        informe('AVISO      pdfplumber no está instalado: no se comprueba la paginación de los PDF.')
    fallidas = []
    try:
        with tempfile.TemporaryDirectory(prefix='stora_fichas_') as tmp:
            tmp = Path(tmp)
            (tmp / CARPETA_FICHAS).mkdir()
            docxs = [escribir_ficha(d, tmp / CARPETA_FICHAS / nombre_ficha(d['id'], 'docx'), fecha, fecha_org)
                     for d, _ in pendientes]
            if docxs:
                informe(f'Convirtiendo {len(docxs)} ficha(s) a PDF…')
            pdfs = convertir_a_pdf(docxs, tmp / CARPETA_FICHAS)
            for (d, h), docx, pdf in zip(pendientes, docxs, pdfs):
                if comprobar:
                    _, problemas = problemas_paginacion(pdf)
                    if problemas:
                        fallidas.append(d['id'])
                        informe(f'ERROR      {nombre_ficha(d["id"], "pdf")}: {"; ".join(problemas)}.')
                        continue
                shutil.copyfile(docx, carpeta / docx.name)
                shutil.copyfile(pdf, carpeta / pdf.name)
                estado['fichas'][d['id']] = {'huella': h, 'generada': fecha}
            indice = escribir_indice(fichas, tmp / f'{INDICE}.docx', fecha)
            pdf_indice, = convertir_a_pdf([indice], tmp, enlaces_relativos=True)
            shutil.copyfile(indice, salida / indice.name)
            shutil.copyfile(pdf_indice, salida / pdf_indice.name)
    except ErrorGeneracion as exc:
        informe(f'ERROR      {exc}')
        return 1
    finally:
        estado['version'] = VERSION
        escribir_json(estado, ruta_estado)

    regeneradas = [d['id'] for d, _ in pendientes if d['id'] not in fallidas]
    informe(f'Fichas: {len(regeneradas)} regenerada(s), {len(fichas) - len(pendientes)} sin cambios '
            f'({len(fichas)} en total) en {carpeta}.')
    if regeneradas:
        informe(f'Regeneradas: {", ".join(regeneradas)}.')
    informe(f'Índice: {salida / (INDICE + ".docx")} y .pdf')
    if fallidas:
        informe(f'{len(fallidas)} ficha(s) con errores de paginación no se han copiado: {", ".join(fallidas)}.')
        return 1
    return 0


# --- validar -------------------------------------------------------------------

def validar(ruta_exp, ruta_org, ruta_panel=None, ruta_tracker=None, salida=None, informe=print):
    errores, avisos = [], []
    exp, errores_exp = cargar_expedientes(ruta_exp)
    errores += errores_exp
    try:
        orgs = cargar_organizaciones(ruta_org)
        filas_panel, avisos_panel = leer_panel(ruta_panel) if ruta_panel else ([], [])
        filas_tracker, avisos_tracker = leer_tracker(ruta_tracker) if ruta_tracker else ([], [])
    except ErrorDatos as exc:
        errores.append(str(exc))
        orgs, filas_panel, filas_tracker, avisos_panel, avisos_tracker = None, [], [], [], []
    avisos += avisos_panel + avisos_tracker
    n = 0
    if exp is not None and not errores_exp and orgs is not None:
        n = len(exp['empresas'])
        _, huerfanos, avisos_cruce = cruzar(exp, orgs, filas_panel, filas_tracker)
        errores += [f'expediente «{oid}» sin correspondencia en {Path(ruta_org).name}.' for oid in huerfanos]
        avisos += [f'«{o["id"]}» ({o.get("nombre", "")}) no tiene expediente: ejecute «inicializar».'
                   for o in orgs if o['id'] not in exp['empresas']]
        avisos += avisos_cruce
    carpeta = Path(salida or Path(ruta_exp).parent) / CARPETA_FICHAS
    pdfs = sorted(carpeta.glob('Ficha_*.pdf')) if carpeta.is_dir() else []
    if pdfs and not hay_pdfplumber():
        avisos.append('pdfplumber no está instalado: no se comprueba la paginación de las fichas generadas.')
    elif pdfs:
        for pdf in pdfs:
            _, problemas = problemas_paginacion(pdf)
            errores += [f'{pdf.name}: {m}.' for m in problemas]
    for m in avisos:
        informe(f'AVISO      {m}')
    for m in errores:
        informe(f'ERROR      {m}')
    informe(f'\n{n} expedientes · {len(pdfs)} fichas generadas comprobadas · {len(errores)} errores · '
            f'{len(avisos)} avisos')
    return 1 if errores else 0


# --- Línea de órdenes ----------------------------------------------------------

def main(argv=None):
    for flujo in (sys.stdout, sys.stderr):
        try:
            flujo.reconfigure(encoding='utf-8')
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='orden', required=True)

    a = sub.add_parser('inicializar', help='crea un expediente vacío para cada organización')
    a.add_argument('--organizaciones', required=True, help='organizaciones.json')
    a.add_argument('--salida', required=True, help='expedientes.json (se crea o se completa)')

    a = sub.add_parser('añadir', aliases=['anadir'], help='aplica un archivo de entradas al expediente')
    a.add_argument('expedientes')
    a.add_argument('entradas')
    a.add_argument('--simular', action='store_true', help='muestra el informe sin escribir nada')
    a.add_argument('--organizaciones', help='permite identificar la empresa por su nombre en organizaciones.json')

    a = sub.add_parser('generar', help='genera las fichas (Word y PDF) y el índice')
    a.add_argument('expedientes')
    a.add_argument('salida', help='carpeta de salida (las fichas van a SALIDA/fichas/)')
    a.add_argument('--organizaciones', required=True)
    a.add_argument('--panel', required=True, help='Seguimiento_Clientes_Stora.xlsx')
    a.add_argument('--tracker', required=True, help='Stora_Exhibition_Lead_Tracker.xlsx')
    a.add_argument('--todas', action='store_true', help='regenera todas las fichas')
    a.add_argument('--fecha', help='fecha de generación (AAAA-MM-DD); por defecto, hoy')
    a.add_argument('--fecha-organizaciones', help='fecha de la copia de organizaciones.json (AAAA-MM-DD)')

    a = sub.add_parser('validar', help='comprueba el expediente, la asociación de nombres y las fichas')
    a.add_argument('expedientes')
    a.add_argument('--organizaciones', required=True)
    a.add_argument('--panel')
    a.add_argument('--tracker')
    a.add_argument('--salida', help='carpeta de salida de «generar»; por defecto, la del expediente')

    args = ap.parse_args(argv)
    for campo in ('fecha', 'fecha_organizaciones'):
        valor = getattr(args, campo, None)
        if valor and not fecha_valida(valor):
            ap.error(f'--{campo.replace("_", "-")} debe tener el formato AAAA-MM-DD')
    if args.orden == 'inicializar':
        return inicializar(args.organizaciones, args.salida)
    if args.orden in ('añadir', 'anadir'):
        return anadir(args.expedientes, args.entradas, args.simular, args.organizaciones)
    if args.orden == 'generar':
        return generar(args.expedientes, args.salida, args.organizaciones, args.panel, args.tracker,
                       args.todas, args.fecha, args.fecha_organizaciones)
    return validar(args.expedientes, args.organizaciones, args.panel, args.tracker, args.salida)


if __name__ == '__main__':
    sys.exit(main())
