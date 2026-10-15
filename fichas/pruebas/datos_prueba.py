"""Excel inventados con la estructura del panel de seguimiento y del tracker de la feria.

Uso:
    python fichas/pruebas/datos_prueba.py [CARPETA]

Crea en CARPETA (por defecto, la de este archivo) Seguimiento_Clientes_Stora.xlsx
y Stora_Exhibition_Lead_Tracker.xlsx con las empresas de
organizaciones_prueba_2026-10-13.json y la que entradas_prueba.json da de alta
fuera de él, más una empresa del panel y otra del tracker que no corresponden a
ninguna. La fila de cabecera no está
en la primera fila, para comprobar que se localiza por su texto.
"""
import datetime
import sys
from pathlib import Path

import openpyxl

AQUI = Path(__file__).resolve().parent
PANEL = 'Seguimiento_Clientes_Stora.xlsx'
TRACKER = 'Stora_Exhibition_Lead_Tracker.xlsx'
SIN_ASOCIAR_PANEL = 'Promotora Desconocida Ejemplo'
SIN_ASOCIAR_TRACKER = 'Empresa Fantasma Tracker'

CABECERA_PANEL = ('Empresa', 'Estado', 'Fecha', 'Contacto', 'Teléfono', 'Email', 'Notas')
FILAS_PANEL = {
    'PROMOTORES': [
        ('Biometano Meseta Ficticia', 'Oferta enviada', datetime.datetime(2026, 10, 20), 'Lucía Ejemplo Sanz',
         '+34 600 000 101', 'lucia.ejemplo@meseta-ficticia.example', 'Prioritaria'),
        (SIN_ASOCIAR_PANEL, 'Primer contacto', None, '', '', '', ''),
    ],
    'EPC': [
        ('Montajes Ejemplo, S.L.', 'Tarifa solicitada', datetime.datetime(2026, 10, 15), 'Andrea Supuesta Ruiz',
         '', '', ''),
        ('Depositos Imaginarios del Levante S.L.', 'Primer contacto', datetime.datetime(2026, 10, 12),
         'Elena Ficticia Mora', '', '', 'Lead nuevo, fuera de organizaciones.json'),
    ],
    'INGENIERÍAS Y CONSULTORÍAS': [
        ('INGENIERIA MODELO NORTE', 'Seguimiento', None, '', '', '', ''),
        ('Consultora Imaginaria', 'Primer contacto', None, '', '', '', ''),
    ],
    'OTROS': [
        ('Cooperativa Agraria Inventada', 'Sin contacto', None, '', '', '', ''),
    ],
}
CABECERA_TRACKER = ('Rating', 'Company', 'Exhibition / Event', 'Contact', 'Opportunity', 'Stage')
FILAS_TRACKER = [  # Rating: entero de 1 a 5 (5 es la máxima), vacío en los leads nuevos
    (5, 'Meseta Bio', 'Feria ficticia de pruebas', 'Lucía Ejemplo Sanz', 'Digestores planta norte', 'Quotation'),
    (4, 'Grupo Montajes Ejemplo', 'Feria ficticia de pruebas', 'Andrea Supuesta Ruiz', 'Tarifa EPC', 'Contacted'),
    (2, 'Consultora Imaginaria', 'Feria ficticia de pruebas', '', '', 'Lead'),
    (None, 'DIL Tanques', 'Feria ficticia de pruebas', 'Elena Ficticia Mora', 'Tanques de agua', 'Lead'),
    (3, SIN_ASOCIAR_TRACKER, 'Feria ficticia de pruebas', '', '', 'Lead'),
]


def crear_panel(ruta):
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for hoja, filas in FILAS_PANEL.items():
        h = libro.create_sheet(hoja)
        h.append(('Seguimiento de clientes Stora (datos inventados)',))
        h.append(())
        h.append(CABECERA_PANEL)
        for f in filas:
            h.append(f)
    libro.save(ruta)
    return ruta


def crear_tracker(ruta):
    libro = openpyxl.Workbook()
    libro.active.title = 'Lead Overview'
    h = libro.active
    h.append(('Exhibition lead tracker (datos inventados)',))
    h.append(CABECERA_TRACKER)
    for f in FILAS_TRACKER:
        h.append(f)
    libro.create_sheet('Notes').append(('Hoja que el generador no lee',))
    libro.save(ruta)
    return ruta


def crear(carpeta=AQUI):
    carpeta = Path(carpeta)
    return crear_panel(carpeta / PANEL), crear_tracker(carpeta / TRACKER)


if __name__ == '__main__':
    for ruta in crear(sys.argv[1] if len(sys.argv) > 1 else AQUI):
        print(f'Creado: {ruta}')
